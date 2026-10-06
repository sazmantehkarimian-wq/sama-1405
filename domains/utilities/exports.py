from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from .models import UtilityAccount, UtilityBill


def _filtered_accounts(request):
    qs = UtilityAccount.objects.select_related("region", "center", "mother_property")
    utility_type = request.GET.get("type", "")
    region = request.GET.get("region", "")
    q = request.GET.get("q", "").strip()
    if utility_type:
        qs = qs.filter(utility_type=utility_type)
    if region:
        qs = qs.filter(region__code=region)
    if q:
        qs = qs.filter(
            Q(title__icontains=q)
            | Q(account_number__icontains=q)
            | Q(meter_number__icontains=q)
            | Q(center__name__icontains=q)
            | Q(mother_property__name__icontains=q)
        )
    return qs


def _style_sheet(ws, max_col):
    thin = Side(style="thin", color="808080")
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A6"
    ws.print_title_rows = "1:5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in ws[5]:
        cell.font = Font(bold=True)


@login_required
def dashboard_xlsx(request):
    accounts = list(_filtered_accounts(request))
    account_ids = [a.id for a in accounts]
    wb = Workbook()
    ws = wb.active
    ws.title = "انشعابات"
    ws.append(["سازمان فرهنگی هنری شهرداری تهران"])
    ws.append(["مدیریت اقتصادی و املاک"])
    ws.append(["اداره املاک و مستغلات"])
    ws.append(["گزارش مدیریتی انشعابات"])
    ws.append(["ردیف", "نوع", "عنوان", "منطقه", "مرکز", "ملک مادر", "شماره اشتراک", "شماره کنتور", "مالکیت", "وضعیت"])
    for i, a in enumerate(accounts, 1):
        ws.append([i, a.get_utility_type_display(), a.title, a.region.name if a.region else "", a.center.name if a.center else "", a.mother_property.name if a.mother_property else "", a.account_number, a.meter_number, a.ownership, a.get_status_display()])
    _style_sheet(ws, 10)
    widths = [8, 14, 28, 18, 28, 32, 22, 20, 18, 18]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = width

    bills = UtilityBill.objects.filter(account_id__in=account_ids).select_related("account", "account__region", "account__center").order_by("-period_end", "-id")
    wbills = wb.create_sheet("قبوض")
    wbills.append(["سازمان فرهنگی هنری شهرداری تهران"])
    wbills.append(["مدیریت اقتصادی و املاک"])
    wbills.append(["اداره املاک و مستغلات"])
    wbills.append(["قبوض مرتبط با گزارش انشعابات"])
    wbills.append(["ردیف", "نوع", "عنوان", "منطقه", "مرکز", "از دوره", "تا دوره", "مبلغ قبض", "سهم سازمان", "سهم بهره‌برداران", "وضعیت"])
    for i, b in enumerate(bills, 1):
        wbills.append([i, b.account.get_utility_type_display(), b.account.title, b.account.region.name if b.account.region else "", b.account.center.name if b.account.center else "", b.period_start, b.period_end, int(b.amount_rial), int(b.organization_amount_rial or 0), int(b.beneficiary_amount_rial or 0), b.get_status_display()])
    _style_sheet(wbills, 11)
    for i, width in enumerate([8, 14, 28, 18, 28, 14, 14, 22, 22, 22, 18], 1):
        wbills.column_dimensions[chr(64+i)].width = width

    stream = BytesIO()
    wb.save(stream)
    response = HttpResponse(stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="utilities-management-report.xlsx"'
    return response
