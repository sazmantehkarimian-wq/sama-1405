from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from .forms import (AccountForm, BillForm, MeterReadingForm, PaymentForm, PolicyForm, ProfileCreateForm, SpaceProfileForm)
from .models import UtilityAccount, UtilityAllocationPolicy, UtilityBill, UtilitySpaceProfile
from .services import calculate_bill


@login_required
def dashboard(request):
    accounts = UtilityAccount.objects.select_related("region", "center", "mother_property").annotate(bill_count=Count("bills"))
    utility_type = request.GET.get("type", "")
    region = request.GET.get("region", "")
    q = request.GET.get("q", "").strip()
    if utility_type:
        accounts = accounts.filter(utility_type=utility_type)
    if region:
        accounts = accounts.filter(region__code=region)
    if q:
        from django.db.models import Q
        accounts = accounts.filter(
            Q(title__icontains=q)
            | Q(account_number__icontains=q)
            | Q(meter_number__icontains=q)
            | Q(center__name__icontains=q)
            | Q(mother_property__name__icontains=q)
        )
    kpis = {
        "accounts": accounts.count(),
        "open_bills": UtilityBill.objects.exclude(status__in=[UtilityBill.Status.PAID, UtilityBill.Status.CANCELLED]).count(),
        "debt_rial": UtilityBill.objects.exclude(status__in=[UtilityBill.Status.PAID, UtilityBill.Status.CANCELLED]).aggregate(v=Sum("amount_rial"))["v"] or 0,
        "calculated": UtilityBill.objects.filter(status=UtilityBill.Status.CALCULATED).count(),
    }
    return render(request, "utilities/dashboard.html", {"accounts": accounts[:200], "kpis": kpis, "types": UtilityAccount.Type.choices})


@login_required
def account_create(request):
    form = AccountForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        account = form.save(commit=False)
        account.created_by = request.user
        account.full_clean()
        account.save()
        UtilityAllocationPolicy.objects.get_or_create(account=account)
        messages.success(request, "انشعاب ایجاد شد.")
        return redirect("utilities:account-detail", pk=account.pk)
    return render(request, "utilities/form.html", {"form": form, "title": "تعریف انشعاب جدید"})


@login_required
def account_edit(request, pk):
    account = get_object_or_404(UtilityAccount, pk=pk)
    form = AccountForm(request.POST or None, instance=account)
    if request.method == "POST" and form.is_valid():
        obj = form.save(commit=False)
        obj.full_clean()
        obj.save()
        messages.success(request, "مشخصات انشعاب ویرایش شد.")
        return redirect("utilities:account-detail", pk=obj.pk)
    return render(request, "utilities/form.html", {"form": form, "title": f"ویرایش انشعاب — {account.title}"})


@login_required
def profile_create(request, account_id):
    account = get_object_or_404(UtilityAccount, pk=account_id)
    form = ProfileCreateForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        profile = form.save(commit=False)
        profile.account = account
        profile.updated_by = request.user
        profile.full_clean()
        profile.save()
        messages.success(request, f"کدفضا {profile.space.code} به انشعاب اضافه شد.")
        return redirect("utilities:account-detail", pk=account.pk)
    return render(request, "utilities/form.html", {"form": form, "title": f"افزودن کدفضا — {account.title}"})


@login_required
def account_detail(request, pk):
    account = get_object_or_404(UtilityAccount.objects.select_related("region", "center", "mother_property"), pk=pk)
    policy, _ = UtilityAllocationPolicy.objects.get_or_create(account=account)
    profiles = account.space_profiles.select_related("space").all()
    bills = account.bills.all()[:50]
    return render(request, "utilities/account_detail.html", {"account": account, "policy": policy, "profiles": profiles, "bills": bills})


@login_required
def policy_edit(request, pk):
    account = get_object_or_404(UtilityAccount, pk=pk)
    policy, _ = UtilityAllocationPolicy.objects.get_or_create(account=account)
    form = PolicyForm(request.POST or None, instance=policy)
    if request.method == "POST" and form.is_valid():
        obj = form.save(commit=False)
        obj.updated_by = request.user
        obj.full_clean()
        obj.save()
        messages.success(request, "سهم سازمان و بهره‌بردار ذخیره شد.")
        return redirect("utilities:account-detail", pk=account.pk)
    return render(request, "utilities/form.html", {"form": form, "title": "ویرایش سهم سازمان / بهره‌بردار — روش سوم"})


@login_required
def profile_edit(request, pk):
    profile = get_object_or_404(UtilitySpaceProfile.objects.select_related("account", "space"), pk=pk)
    form = SpaceProfileForm(request.POST or None, instance=profile)
    if request.method == "POST" and form.is_valid():
        obj = form.save(commit=False)
        obj.updated_by = request.user
        obj.full_clean()
        obj.save()
        messages.success(request, f"پارامترهای کدفضا {obj.space.code} ذخیره شد.")
        return redirect("utilities:account-detail", pk=obj.account_id)
    return render(request, "utilities/form.html", {"form": form, "title": f"تنظیم روش سوم — کدفضا {profile.space.code}"})


@login_required
def bill_create(request, account_id):
    account = get_object_or_404(UtilityAccount, pk=account_id)
    form = BillForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        bill = form.save(commit=False)
        bill.account = account
        bill.created_by = request.user
        bill.full_clean()
        bill.save()
        messages.success(request, "قبض ثبت شد؛ اکنون می‌توانید محاسبه روش سوم را اجرا کنید.")
        return redirect("utilities:bill-detail", pk=bill.pk)
    return render(request, "utilities/form.html", {"form": form, "title": f"ثبت قبض جدید — {account}"})


@login_required
def bill_detail(request, pk):
    bill = get_object_or_404(UtilityBill.objects.select_related("account"), pk=pk)
    allocations = bill.allocations.select_related("profile__space").all()
    return render(request, "utilities/bill_detail.html", {"bill": bill, "allocations": allocations})


@login_required
def bill_calculate(request, pk):
    bill = get_object_or_404(UtilityBill, pk=pk)
    if request.method != "POST":
        return redirect("utilities:bill-detail", pk=pk)
    try:
        calculate_bill(bill, persist=True)
    except (ValueError, AssertionError) as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, "محاسبه روش سوم انجام شد و کنترل جمع قبض برقرار است.")
    return redirect("utilities:bill-detail", pk=pk)


@login_required
def reading_create(request, bill_id):
    bill = get_object_or_404(UtilityBill, pk=bill_id)
    form = MeterReadingForm(request.POST or None, bill=bill)
    if request.method == "POST" and form.is_valid():
        reading = form.save(commit=False)
        reading.bill = bill
        reading.full_clean()
        reading.save()
        messages.success(request, "قرائت زیرکنتور ثبت شد. برای اعمال آن، قبض را دوباره محاسبه کنید.")
        return redirect("utilities:bill-detail", pk=bill.pk)
    return render(request, "utilities/form.html", {"form": form, "title": f"ثبت زیرکنتور — قبض {bill.pk}"})


@login_required
def payment_create(request, bill_id):
    bill = get_object_or_404(UtilityBill, pk=bill_id)
    form = PaymentForm(request.POST or None, request.FILES or None, bill=bill)
    if request.method == "POST" and form.is_valid():
        payment = form.save(commit=False)
        payment.bill = bill
        payment.created_by = request.user
        payment.full_clean()
        payment.save()
        paid = bill.payments.aggregate(v=Sum("amount_rial"))["v"] or 0
        if paid >= bill.amount_rial:
            bill.status = UtilityBill.Status.PAID
        elif paid > 0:
            bill.status = UtilityBill.Status.PARTIAL
        bill.save(update_fields=["status", "updated_at"])
        messages.success(request, "پرداخت ثبت شد.")
        return redirect("utilities:bill-detail", pk=bill.pk)
    return render(request, "utilities/form.html", {"form": form, "title": f"ثبت پرداخت — قبض {bill.pk}"})


@login_required
def bill_xlsx(request, pk):
    bill = get_object_or_404(UtilityBill.objects.select_related("account"), pk=pk)
    allocations = bill.allocations.select_related("profile__space").all()
    wb = Workbook()
    ws = wb.active
    ws.title = "گزارش انشعابات"
    ws.sheet_view.rightToLeft = True
    ws.append(["سازمان فرهنگی هنری شهرداری تهران"])
    ws.append(["مدیریت اقتصادی و املاک — اداره املاک و مستغلات"])
    ws.append([f"گزارش قبض {bill.account.get_utility_type_display()} — {bill.account.title}"])
    ws.append([])
    headers = ["ردیف", "کد فضا", "نام فضا", "مبنای محاسبه", "وزن مصرف", "سهم (%)", "مبلغ (ریال)"]
    ws.append(headers)
    for i, row in enumerate(allocations, start=1):
        ws.append([i, row.profile.space.code, row.profile.space.name, row.get_basis_display(), float(row.consumption_weight), float(row.share_percent), int(row.amount_rial)])
    ws.append([])
    ws.append(["سهم سازمان", "", "", "", "", float(bill.organization_percent_effective or 0), int(bill.organization_amount_rial or 0)])
    ws.append(["سهم بهره‌برداران", "", "", "", "", float(bill.beneficiary_percent_effective or 0), int(bill.beneficiary_amount_rial or 0)])
    ws.append(["جمع کنترل", "", "", "", "", "", int((bill.organization_amount_rial or 0) + (bill.beneficiary_amount_rial or 0))])
    thin = Side(style="thin", color="808080")
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=7):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in ws[5]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A6"
    ws.print_title_rows = "1:5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    widths = [8, 12, 28, 28, 18, 16, 22]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    stream = BytesIO()
    wb.save(stream)
    response = HttpResponse(stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="utility-bill-{bill.pk}.xlsx"'
    return response
