import hashlib
import io
import json
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.utils import timezone
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from domains.auctionflow.models import AuctionDocumentInstance, AuctionLotProfile, AuctionPeriodProfile
from domains.contracts.models import Beneficiary
from domains.identity.models import AuditEvent
from domains.operations.models import AuctionLot, AuctionProposal
from services.contract_circulation import create_circulation
from services.documents import store_document

TEMPLATE_VERSION = 'reference-1405-07-12-v1'
ORG_HEADERS = ['سازمان فرهنگی هنری شهرداری تهران', 'مدیریت اقتصادی و املاک', 'اداره املاک و مستغلات']


def _canonical(payload):
    return json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(',', ':'), default=str).encode('utf-8')


def snapshot_hash(payload):
    return hashlib.sha256(_canonical(payload)).hexdigest()


def _money(value):
    if value in (None, ''):
        return '—'
    return f'{Decimal(value):,.0f}'


def ensure_profiles(period):
    period_profile, _ = AuctionPeriodProfile.objects.get_or_create(period=period)
    lot_profiles = []
    for lot in period.lots.select_related('space').all():
        profile, _ = AuctionLotProfile.objects.get_or_create(lot=lot)
        lot_profiles.append(profile)
    return period_profile, lot_profiles


def build_lot_snapshot(lot):
    period_profile, _ = AuctionPeriodProfile.objects.get_or_create(period=lot.period)
    profile, _ = AuctionLotProfile.objects.get_or_create(lot=lot)
    space = lot.space
    proposals = []
    for proposal in lot.proposals.select_related('participant').order_by('pk'):
        proposals.append({
            'proposal_id': proposal.pk,
            'participant_id': proposal.participant_id,
            'participant_name': proposal.participant.name,
            'identity_number': proposal.participant.identity_number,
            'contact': proposal.participant.contact,
            'envelope_a': proposal.envelope_a_received,
            'envelope_b': proposal.envelope_b_received,
            'envelope_c': proposal.envelope_c_received,
            'offered_amount_rial': str(proposal.offered_amount_rial) if proposal.offered_amount_rial is not None else None,
            'status': proposal.status,
        })
    current_appraisal = space.appraisals.filter(is_current=True).order_by('-pk').first()
    return {
        'organization': ORG_HEADERS,
        'period': {
            'id': lot.period_id, 'identity': lot.period.identity, 'title': lot.period.title,
            'planned_date': lot.period.planned_date, 'state': lot.period.state,
            'permit_reference': period_profile.permit_reference,
            'ad_day_name': period_profile.ad_day_name, 'ad_date': period_profile.ad_date,
            'newspaper': period_profile.newspaper, 'invitation_number': period_profile.invitation_number,
            'opening_session_date': period_profile.opening_session_date,
            'opening_session_time': period_profile.opening_session_time,
            'opening_session_location': period_profile.opening_session_location,
            'duration_years': period_profile.duration_years,
        },
        'lot': {
            'id': lot.pk, 'space_code': space.code, 'space_name': space.name,
            'area': str(space.area) if space.area is not None else None,
            'address': space.address, 'usage': space.current_usage,
            'proposed_activity': space.proposed_activity, 'activity_group': space.activity_group,
            'entry_method': lot.entry_method, 'readiness': lot.readiness,
            'template_family': profile.template_family,
            'base_monthly_rent_rial': str(profile.base_monthly_rent_rial) if profile.base_monthly_rent_rial is not None else None,
            'guarantee_amount_rial': str(profile.guarantee_amount_rial) if profile.guarantee_amount_rial is not None else None,
            'investment_amount_rial': str(profile.investment_amount_rial) if profile.investment_amount_rial is not None else None,
            'proposed_job': profile.proposed_job,
            'decision_reference': profile.decision_reference, 'decision_date': profile.decision_date,
            'result_state': profile.result_state,
        },
        'current_appraisal': {
            'id': current_appraisal.pk if current_appraisal else None,
            'amount_rial': str(current_appraisal.amount_rial) if current_appraisal and current_appraisal.amount_rial is not None else None,
            'date': current_appraisal.appraisal_date if current_appraisal else None,
            'reference': current_appraisal.reference if current_appraisal else None,
        },
        'proposals': proposals,
        'winner_proposal_id': profile.winner_proposal_id,
        'runner_up_proposal_id': profile.runner_up_proposal_id,
        'beneficiary_id': profile.beneficiary_id,
        'contract_circulation_id': profile.contract_circulation_id,
        'official_contract_id': profile.official_contract_id,
    }


def validate_winner_proposal(lot, proposal):
    if proposal.lot_id != lot.pk:
        raise ValidationError('پیشنهاد انتخاب‌شده متعلق به این فضای مزایده نیست.')
    missing = []
    if not proposal.envelope_a_received: missing.append('پاکت الف')
    if not proposal.envelope_b_received: missing.append('پاکت ب')
    if not proposal.envelope_c_received: missing.append('پاکت ج')
    if proposal.offered_amount_rial is None: missing.append('مبلغ پیشنهاد')
    if missing:
        raise ValidationError('تعیین برنده تا تکمیل این موارد مجاز نیست: ' + '، '.join(missing))


@transaction.atomic
def select_winner(*, lot, winner_proposal, runner_up_proposal=None, decision_reference, decision_date, actor, ip_address=None):
    profile, _ = AuctionLotProfile.objects.select_for_update().get_or_create(lot=lot)
    validate_winner_proposal(lot, winner_proposal)
    if runner_up_proposal:
        validate_winner_proposal(lot, runner_up_proposal)
        if runner_up_proposal.pk == winner_proposal.pk:
            raise ValidationError('برنده و نفر دوم نمی‌توانند یک پیشنهاد باشند.')
    if not (decision_reference or '').strip():
        raise ValidationError('مرجع تصمیم کمیسیون برای تعیین برنده الزامی است.')
    if not (decision_date or '').strip():
        raise ValidationError('تاریخ تصمیم کمیسیون برای تعیین برنده الزامی است.')
    profile.winner_proposal = winner_proposal
    profile.runner_up_proposal = runner_up_proposal
    profile.decision_reference = decision_reference.strip()
    profile.decision_date = decision_date.strip()
    profile.result_state = AuctionLotProfile.ResultState.AWARDED
    profile.awarded_by = actor
    profile.awarded_at = timezone.now()
    lot.winner_name = winner_proposal.participant.name
    lot.winning_amount_rial = winner_proposal.offered_amount_rial
    lot.result = 'برنده تعیین شد'
    lot.save(update_fields=['winner_name', 'winning_amount_rial', 'result'])
    profile.snapshot = build_lot_snapshot(lot)
    profile.save()
    AuditEvent.objects.create(
        actor=actor, action='AUCTION_WINNER_SELECTED', entity_type='AuctionLot', entity_id=str(lot.pk),
        after={'winner_proposal_id': winner_proposal.pk, 'runner_up_proposal_id': runner_up_proposal.pk if runner_up_proposal else None,
               'winner': winner_proposal.participant.name, 'amount_rial': str(winner_proposal.offered_amount_rial),
               'decision_reference': profile.decision_reference, 'decision_date': profile.decision_date},
        ip_address=ip_address,
    )
    return profile


def _split_contact(contact):
    text = (contact or '').strip()
    return text if len(text) <= 20 else ''


@transaction.atomic
def start_contract_from_award(*, lot, actor, beneficiary_kind, operational_start_date, due_date='', ip_address=None):
    profile, _ = AuctionLotProfile.objects.select_for_update().get_or_create(lot=lot)
    if not profile.winner_proposal_id:
        raise ValidationError('ابتدا برنده مزایده را ثبت کنید.')
    if profile.contract_circulation_id:
        return profile.contract_circulation
    proposal = profile.winner_proposal
    validate_winner_proposal(lot, proposal)
    participant = proposal.participant
    beneficiary = None
    if participant.identity_number:
        beneficiary = Beneficiary.objects.filter(identity_number=participant.identity_number, archived_at__isnull=True).first()
    if not beneficiary:
        beneficiary = Beneficiary.objects.create(
            name=participant.name,
            identity_number=participant.identity_number,
            kind=beneficiary_kind,
            mobile=_split_contact(participant.contact),
            contact=participant.contact,
            created_by=actor,
        )
    circulation = create_circulation(
        space=lot.space, beneficiary=beneficiary,
        subject=f'انعقاد قرارداد برنده مزایده {lot.period.identity} — فضای {lot.space.code}',
        operational_start_date=operational_start_date,
        next_action='تکمیل تضامین، کنترل اسناد و گردش امضای قرارداد برنده مزایده',
        due_date=due_date, actor=actor, ip_address=ip_address,
    )
    profile.beneficiary = beneficiary
    profile.contract_circulation = circulation
    profile.result_state = AuctionLotProfile.ResultState.CONTRACTING
    profile.snapshot = build_lot_snapshot(lot)
    profile.save(update_fields=['beneficiary','contract_circulation','result_state','snapshot','updated_at'])
    AuditEvent.objects.create(
        actor=actor, action='AUCTION_CONTRACT_FLOW_STARTED', entity_type='AuctionLot', entity_id=str(lot.pk),
        after={'beneficiary_id': beneficiary.pk, 'contract_circulation_id': circulation.pk}, ip_address=ip_address,
    )
    return circulation


@transaction.atomic
def sync_official_contract(*, circulation, actor=None, ip_address=None):
    try:
        profile = circulation.auction_award
    except AuctionLotProfile.DoesNotExist:
        return None
    if not circulation.official_contract_id:
        return profile
    if profile.official_contract_id != circulation.official_contract_id or profile.result_state != AuctionLotProfile.ResultState.CONTRACTED:
        profile.official_contract = circulation.official_contract
        profile.result_state = AuctionLotProfile.ResultState.CONTRACTED
        profile.snapshot = build_lot_snapshot(profile.lot)
        profile.save(update_fields=['official_contract','result_state','snapshot','updated_at'])
        if actor:
            AuditEvent.objects.create(actor=actor, action='AUCTION_CONTRACT_FINALIZED', entity_type='AuctionLot', entity_id=str(profile.lot_id), after={'contract_id': circulation.official_contract_id}, ip_address=ip_address)
    return profile


def _rtl(paragraph, bold=False):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1'); ppr.append(bidi)
    for run in paragraph.runs:
        run.font.name = 'Vazirmatn'; run.bold = bold
        if run._element.rPr is not None:
            run._element.rPr.rFonts.set(qn('w:cs'), 'Vazirmatn')


def _add_field(doc, label, value):
    p = doc.add_paragraph()
    r = p.add_run(f'{label}: '); r.bold = True
    p.add_run(str(value if value not in (None, '') else '—'))
    _rtl(p)


def _docx_from_snapshot(document_type, snapshot):
    doc = DocxDocument()
    for header in ORG_HEADERS:
        p=doc.add_paragraph(header); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; _rtl(p, True)
    p=doc.add_paragraph('پیش‌نویس کنترل‌شده UAT — مبتنی بر نمونه‌های مرجع مالک؛ تا تأیید چاپ Golden Master سند رسمی نهایی نیست.')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER; _rtl(p, True)
    lot=snapshot['lot']; period=snapshot['period']; appraisal=snapshot['current_appraisal']
    titles={
        'CONDITIONS':'شرایط عمومی و اختصاصی مزایده',
        'PRICE_FORM':'فرم پیشنهاد قیمت / روکش پاکات',
        'OPENING_MINUTES':'صورتجلسه بازگشایی پاکات',
        'EXPERT_NOTICE':'ابلاغ ارزیابی کارشناس رسمی',
        'SAMPLE_CONTRACT':'نمونه قرارداد برنده مزایده',
    }
    h=doc.add_heading(titles.get(document_type, document_type), level=1); h.alignment=WD_ALIGN_PARAGRAPH.CENTER; _rtl(h, True)
    _add_field(doc,'شناسه دوره',period['identity']); _add_field(doc,'عنوان دوره',period['title'])
    _add_field(doc,'کد فضا',lot['space_code']); _add_field(doc,'نام فضا',lot['space_name'])
    _add_field(doc,'کاربری',lot['usage']); _add_field(doc,'مساحت',lot['area']); _add_field(doc,'نشانی',lot['address'])
    if document_type in {'CONDITIONS','PRICE_FORM','OPENING_MINUTES','SAMPLE_CONTRACT'}:
        _add_field(doc,'اجاره پایه ماهانه (ریال)',_money(lot['base_monthly_rent_rial']))
        _add_field(doc,'تضمین شرکت در مزایده (ریال)',_money(lot['guarantee_amount_rial']))
        _add_field(doc,'سرمایه‌گذاری (ریال)',_money(lot['investment_amount_rial']))
    if document_type=='CONDITIONS':
        _add_field(doc,'روز درج آگهی',period['ad_day_name']); _add_field(doc,'تاریخ درج آگهی',period['ad_date']); _add_field(doc,'روزنامه',period['newspaper'])
        p=doc.add_paragraph('متن ثابت شرایط عمومی و اختصاصی باید عیناً از فایل مرجع همان خانواده منتقل شود. این نسخه UAT فقط Field Map و داده‌های تکمیل‌شونده را تثبیت می‌کند و جایگزین Golden Master تأییدشده نیست.'); _rtl(p)
    elif document_type=='PRICE_FORM':
        p=doc.add_paragraph('نام پیشنهاددهنده: ........................................     مبلغ پیشنهادی ماهانه: ........................................ ریال'); _rtl(p)
        p=doc.add_paragraph('مدارک پاکت‌ها مطابق نمونه مرجع توسط پیشنهاددهنده تکمیل و امضا می‌شود.'); _rtl(p)
    elif document_type=='OPENING_MINUTES':
        _add_field(doc,'مجوز برگزاری مزایده',period['permit_reference']); _add_field(doc,'مدت واگذاری',f"{period['duration_years']} سال شمسی")
        _add_field(doc,'شماره دعوتنامه',period['invitation_number']); _add_field(doc,'تاریخ جلسه',period['opening_session_date']); _add_field(doc,'ساعت',period['opening_session_time']); _add_field(doc,'محل جلسه',period['opening_session_location'])
        table=doc.add_table(rows=1, cols=6); table.style='Table Grid'
        headers=['نام شرکت‌کننده','پاکت الف','پاکت ب','پاکت ج','مبلغ پیشنهادی','وضعیت']
        for i,x in enumerate(headers): table.rows[0].cells[i].text=x
        for proposal in snapshot['proposals']:
            cells=table.add_row().cells
            vals=[proposal['participant_name'],'بله' if proposal['envelope_a'] else 'خیر','بله' if proposal['envelope_b'] else 'خیر','بله' if proposal['envelope_c'] else 'خیر',_money(proposal['offered_amount_rial']),proposal['status'] or '—']
            for i,v in enumerate(vals): cells[i].text=str(v)
        if snapshot.get('winner_proposal_id'):
            winner=next((x for x in snapshot['proposals'] if x['proposal_id']==snapshot['winner_proposal_id']),None)
            if winner:
                _add_field(doc,'برنده ثبت‌شده بر اساس تصمیم کمیسیون',winner['participant_name']); _add_field(doc,'مبلغ برنده',_money(winner['offered_amount_rial']))
                _add_field(doc,'مرجع تصمیم',lot['decision_reference']); _add_field(doc,'تاریخ تصمیم',lot['decision_date'])
    elif document_type=='EXPERT_NOTICE':
        p=doc.add_paragraph('احتراماً، خواهشمند است ضمن بازدید و ارزیابی فضای زیر نسبت به تعیین پایه اجاره‌بهای ماهانه اقدام و نتیجه را به این مدیریت گزارش فرمایید.'); _rtl(p)
        _add_field(doc,'عنوان مرکز و کاربری',f"{lot['space_name']} — {lot['usage'] or lot['proposed_activity']}")
        _add_field(doc,'رقم کارشناسی جاری',_money(appraisal['amount_rial'])); _add_field(doc,'مرجع کارشناسی جاری',appraisal['reference'])
    elif document_type=='SAMPLE_CONTRACT':
        winner=next((x for x in snapshot['proposals'] if x['proposal_id']==snapshot.get('winner_proposal_id')),None)
        _add_field(doc,'طرف منتخب مزایده',winner['participant_name'] if winner else 'پس از تعیین برنده تکمیل می‌شود')
        _add_field(doc,'شناسه هویتی',winner['identity_number'] if winner else '')
        _add_field(doc,'مبلغ اجاره ماهانه برنده',_money(winner['offered_amount_rial']) if winner else '—')
        _add_field(doc,'مدت',f"{period['duration_years']} سال شمسی")
        p=doc.add_paragraph('متن حقوقی ثابت قرارداد باید از نمونه PDF خانواده تجاری/کافه/ورزشی بدون دخل و تصرف به Golden Master منتقل شود. این خروجی صرفاً پیش‌نویس داده‌ای UAT برای کنترل Field Mapping است.'); _rtl(p)
    out=io.BytesIO(); doc.save(out); return out.getvalue()


@transaction.atomic
def generate_controlled_document(*, lot, document_type, actor, ip_address=None):
    profile, _ = AuctionLotProfile.objects.get_or_create(lot=lot)
    period_profile, _ = AuctionPeriodProfile.objects.get_or_create(period=lot.period)
    if document_type == 'SAMPLE_CONTRACT' and not profile.winner_proposal_id:
        raise ValidationError('نمونه قرارداد تکمیل‌شده فقط پس از ثبت برنده قابل تولید است.')
    snapshot=build_lot_snapshot(lot)
    snap_hash=snapshot_hash(snapshot)
    payload=_docx_from_snapshot(document_type, snapshot)
    source_identity=f'OWNER_REFERENCE_BUNDLE::{profile.template_family}::{document_type}::{period_profile.template_version}'
    source_sha=hashlib.sha256(source_identity.encode('utf-8')).hexdigest()
    filename=f"auction-{lot.period.identity}-{lot.space.code}-{document_type.lower()}.docx"
    uploaded=SimpleUploadedFile(filename,payload,content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    stored=store_document(
        uploaded=uploaded, title=f'{AuctionDocumentInstance.DocumentType(document_type).label} — فضای {lot.space.code}',
        document_type=f'AUCTION_GENERATED_{document_type}', entity_type='AuctionLot', entity_id=lot.pk,
        user=actor, reference=lot.period.identity, document_date=lot.period.planned_date,
        notes=f'پیش‌نویس کنترل‌شده UAT؛ template={period_profile.template_version}; snapshot={snap_hash}', ip_address=ip_address,
    )
    instance=AuctionDocumentInstance.objects.create(
        period=lot.period, lot=lot, document_type=document_type, template_family=profile.template_family,
        template_version=period_profile.template_version, source_sha256=source_sha,
        snapshot=snapshot, snapshot_sha256=snap_hash, document=stored, generated_by=actor,
    )
    AuditEvent.objects.create(actor=actor,action='AUCTION_DOCUMENT_GENERATED',entity_type='AuctionLot',entity_id=str(lot.pk),after={'instance_id':instance.pk,'document_type':document_type,'template_version':instance.template_version,'snapshot_sha256':snap_hash},ip_address=ip_address)
    return instance
