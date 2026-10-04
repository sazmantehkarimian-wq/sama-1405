import hashlib
import io
import os
import zipfile

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from docx import Document as DocxDocument

from domains.auctionflow.models import AuctionTemplateSource
from services.documents import store_document


REFERENCE_VERSION = 'reference-1405-07-12-v1'
REFERENCE_FILES = {
    'نمونه رزوکش پاکات.docx': ('GENERAL', 'ENVELOPE_COVER', 'dff758d7549f1effc59733b040d91da7eb67032e9c22d4f16e1bf9833b4c281f'),
    'شرایط عمومی اختصاصی کافه.docx': ('CAFE', 'CONDITIONS', '0926059b336ba39e76fa0186a67085d62bcd09f01be311c724a527c8267b2381'),
    'نمونه قرارداد کافه.pdf': ('CAFE', 'SAMPLE_CONTRACT', '5d4ba166e3cda354702e1782dc19a93c9c8d22806582c2fc982e429e5a424690'),
    'شرایط عمومی اختصاصی فضای تجاری.docx': ('COMMERCIAL', 'CONDITIONS', '925d4e36d5d63c81645e2c67c1a356b7e7fc89e167d6f98ff55e09d19a369131'),
    'فرم نمونه ابلاغ کارشناس رسمی.docx': ('GENERAL', 'EXPERT_NOTICE', '13adc308ff902dda352fe59ad90490557744e7ff646ca606db70612b77486860'),
    'نمونه صورتجلسه بازگشایی پاکات.docx': ('GENERAL', 'OPENING_MINUTES', '2d698909a182c249a218bd75af2e770d5be91156d7d6c84f3715860e3b77b566'),
    'شرایط عمومی اختصاصی فضای ورزشی.docx': ('SPORT', 'CONDITIONS', '6b2dc335a16767eabcfcc91c30dd5417efd4dbf5170e5d744982000e904f4553'),
    'نمونه قرارداد تجاری.pdf': ('COMMERCIAL', 'SAMPLE_CONTRACT', '811ac0e9bc853dae9c7d172493c038dc9007a61a08f495ebf19892de06f66a9b'),
    'نمونه قرارداد ورزشی.pdf': ('SPORT', 'SAMPLE_CONTRACT', '15b61ebf77d4f123647b326d7cd117ec4176c95a8dd0f2332f89846d8b47fc2c'),
}


def _basename(name):
    return os.path.basename(name.replace('\\', '/'))


def _sha(data):
    return hashlib.sha256(data).hexdigest()


@transaction.atomic
def import_reference_pack(*, uploaded, actor, ip_address=None):
    raw = uploaded.read()
    if len(raw) > 20 * 1024 * 1024:
        raise ValidationError('حجم بسته قالب‌های مزایده بیش از حد مجاز است.')
    try:
        archive = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile as exc:
        raise ValidationError('فایل انتخاب‌شده ZIP معتبر نیست.') from exc

    found = {}
    for info in archive.infolist():
        if info.is_dir():
            continue
        if info.file_size > 10 * 1024 * 1024:
            raise ValidationError('یکی از فایل‌های قالب بیش از حد مجاز است.')
        name = _basename(info.filename)
        if name in REFERENCE_FILES:
            if name in found:
                raise ValidationError(f'فایل تکراری در بسته: {name}')
            found[name] = archive.read(info)

    missing = [name for name in REFERENCE_FILES if name not in found]
    if missing:
        raise ValidationError('بسته ناقص است. فایل‌های مفقود: ' + '، '.join(missing))

    mismatches = []
    for name, data in found.items():
        expected = REFERENCE_FILES[name][2]
        if _sha(data) != expected:
            mismatches.append(name)
    if mismatches:
        raise ValidationError('هش فایل مرجع با نسخه مصوب مالک تطابق ندارد: ' + '، '.join(mismatches))

    created = []
    for name, data in found.items():
        family, kind, expected_sha = REFERENCE_FILES[name]
        AuctionTemplateSource.objects.filter(family=family, kind=kind, active=True).update(active=False)
        mime = 'application/pdf' if name.lower().endswith('.pdf') else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        document = store_document(
            uploaded=SimpleUploadedFile(name, data, content_type=mime),
            title=f'قالب مرجع مزایده — {name}',
            document_type='AUCTION_TEMPLATE_SOURCE',
            entity_type='AuctionTemplateSource',
            entity_id=f'{family}:{kind}:{REFERENCE_VERSION}',
            user=actor,
            reference=REFERENCE_VERSION,
            notes=f'Owner Reference Template; SHA-256={expected_sha}; تا تأیید چاپی Golden Master رسمی نیست.',
            ip_address=ip_address,
        )
        source = AuctionTemplateSource.objects.create(
            family=family, kind=kind, version=REFERENCE_VERSION,
            source_filename=name, source_sha256=expected_sha, source_document=document,
            active=True, imported_by=actor,
        )
        created.append(source)
    return created


def source_for(*, document_type, family):
    if document_type == 'CONDITIONS':
        key_family, kind = family, 'CONDITIONS'
    elif document_type in {'PRICE_FORM', 'ENVELOPE_A', 'ENVELOPE_B', 'ENVELOPE_C'}:
        key_family, kind = 'GENERAL', 'ENVELOPE_COVER'
    elif document_type == 'OPENING_MINUTES':
        key_family, kind = 'GENERAL', 'OPENING_MINUTES'
    elif document_type == 'EXPERT_NOTICE':
        key_family, kind = 'GENERAL', 'EXPERT_NOTICE'
    elif document_type == 'SAMPLE_CONTRACT':
        key_family, kind = family, 'SAMPLE_CONTRACT'
    else:
        raise ValidationError('نوع سند برای موتور قالب مرجع تعریف نشده است.')
    source = AuctionTemplateSource.objects.filter(family=key_family, kind=kind, active=True).select_related('source_document').order_by('-imported_at').first()
    if not source:
        raise ValidationError('ابتدا بسته قالب‌های مرجع مزایده را یک‌بار در تنظیمات مزایده ثبت کنید.')
    return source


def source_bytes(source):
    with source.source_document.file.open('rb') as handle:
        data = handle.read()
    digest = _sha(data)
    if digest != source.source_sha256:
        raise ValidationError('فایل Master تغییر کرده است؛ تولید سند متوقف شد.')
    return data


def _set_run(paragraph, run_index, value):
    if run_index >= len(paragraph.runs):
        raise ValidationError('ساختار قالب مرجع با Field Map سازگار نیست.')
    paragraph.runs[run_index].text = str(value or '')


def _fill_common_cover(doc, fields):
    # Exact paragraph positions are safe because source SHA is pinned.
    _set_run(doc.paragraphs[2], 0, f"مزایده عمومی واگذاری و بهره برداری و اجاره {fields.get('SPACE-002','')}")
    _set_run(doc.paragraphs[3], 0, f"واقع در {fields.get('SPACE-005','')}")
    _set_run(doc.paragraphs[4], 0, f"با کاربری {fields.get('SPACE-004','')} کد {fields.get('SPACE-001','')}")


def _fill_conditions(doc, fields):
    p = doc.paragraphs
    # Page 2: advertisement date/day, office hours. Participant blanks remain untouched.
    _set_run(p[10], 1, fields.get('AUC-013') or fields.get('AUC-006') or '')
    _set_run(p[10], 3, (fields.get('AUC-006') or '') + ' و')
    if fields.get('ORG-011'):
        # Keep the static label and only replace the variable office-hours phrase.
        _set_run(p[10], 10, ' ')
        _set_run(p[10], 11, fields['ORG-011'])
        for idx in (12,13,14,15,16):
            _set_run(p[10], idx, '')
    # Main lot description and approved variable values.
    _set_run(p[68], 1, fields.get('SPACE-002',''))
    _set_run(p[68], 3, f"{fields.get('SPACE-005','')} با کد {fields.get('SPACE-001','')}")
    _set_run(p[68], 8, fields.get('SPACE-004',''))
    _set_run(p[70], 6, fields.get('LOT-005',''))
    _set_run(p[98], 1, fields.get('LOT-003',''))
    _set_run(p[118], 4, fields.get('AUC-013',''))
    _set_run(p[118], 7, fields.get('AUC-012',''))
    if fields.get('ORG-012'):
        _set_run(p[118], 10, fields['ORG-012'])
    _set_run(p[121], 2, fields.get('AUC-015',''))
    _set_run(p[121], 4, fields.get('AUC-014',''))
    if fields.get('AUC-017'):
        _set_run(p[121], 8, fields['AUC-017'])
        _set_run(p[121], 9, '')
    # Embedded distribution forms: only organization/space fields are auto-filled.
    for start in (130, 174, 206):
        p[start].text = f"مزایده عمومی واگذاری و بهره برداری و اجاره {fields.get('SPACE-002','')}"
        p[start+1].text = f"واقع در {fields.get('SPACE-005','')}"
        p[start+2].text = f"با کاربری {fields.get('SPACE-004','')} کد {fields.get('SPACE-001','')}"
    if doc.tables:
        table = doc.tables[0]
        table.rows[1].cells[0].text = fields.get('SPACE-002','')
        table.rows[1].cells[1].text = fields.get('SPACE-006','')
        table.rows[1].cells[2].text = fields.get('LOT-003','')
        table.rows[1].cells[3].text = fields.get('LOT-001','')
        table.rows[2].cells[0].text = f"{fields.get('SPACE-005','')} — هماهنگی: {fields.get('SPACE-009','')} {fields.get('SPACE-008','')}".strip()


def _fill_expert_notice(doc, fields):
    doc.paragraphs[2].text = f"عنوان مركز و كاربري : {fields.get('SPACE-002','')} — {fields.get('SPACE-004','')}"
    doc.paragraphs[3].text = f"-تعيين اجاره بهاي فضای کد {fields.get('SPACE-001','')} جهت {fields.get('SPACE-004','')}"
    if fields.get('SPACE-005'):
        doc.paragraphs[5].text = fields['SPACE-005']


def _fill_opening_minutes(doc, fields):
    table = doc.tables[0]
    mapping = {
        0: 'AUC-004', 1: 'LOT-005', 2: 'LOT-001', 3: 'LOT-007', 4: 'LOT-003',
        5: 'AUC-006', 6: 'AUC-018', 7: 'AUC-014', 8: 'AUC-017', 9: 'SESSION-004',
    }
    for row, field_id in mapping.items():
        if fields.get(field_id) not in (None, ''):
            table.rows[row].cells[1].text = str(fields[field_id])
    # Session participant/proposal tables are filled only at opening-session stage.
    participants = fields.get('_SESSION_PROPOSALS') or []
    if participants:
        participant_table = doc.tables[1]
        while len(participant_table.rows) < len(participants) + 1:
            participant_table.add_row()
        price_table = doc.tables[3]
        while len(price_table.rows) < len(participants) + 1:
            price_table.add_row()
        for idx, item in enumerate(participants, start=1):
            vals = [item.get('name',''), item.get('contact',''), item.get('postal_code',''), item.get('birth_date',''), item.get('identity_number',''), item.get('receipt_number',''), item.get('guarantee_amount','')]
            for ci, value in enumerate(vals):
                participant_table.rows[idx].cells[ci].text = str(value or '')
            price_table.rows[idx].cells[0].text = str(idx)
            price_table.rows[idx].cells[1].text = str(item.get('name',''))
            price_table.rows[idx].cells[2].text = str(item.get('offered_amount',''))
    winner = fields.get('_WINNER')
    if winner:
        doc.paragraphs[28].text = f"نتیجه مزایده: شغل پیشنهادی: {fields.get('PART-017','')}"
        doc.paragraphs[29].text = (
            f"اعضاء با عنایت به صرفه وصلاح سازمان {winner.get('name','')} را با پیشنهاد قیمت ماهانه به مبلغ "
            f"{winner.get('amount','')} ریال بعنوان برنده مزایده اعلام نمودند و مقرر گردید طبق شرایط عمومی و اختصاصی مزایده "
            f"با نامبرده از زمان تحویل فضا به مدت {fields.get('LOT-005','')} سال شمسی، قرارداد منعقد گردد."
        )


def render_docx_from_source(*, source, document_type, fields):
    data = source_bytes(source)
    doc = DocxDocument(io.BytesIO(data))
    if document_type == 'CONDITIONS':
        _fill_conditions(doc, fields)
    elif document_type in {'PRICE_FORM', 'ENVELOPE_A', 'ENVELOPE_B', 'ENVELOPE_C'}:
        _fill_common_cover(doc, fields)
    elif document_type == 'EXPERT_NOTICE':
        _fill_expert_notice(doc, fields)
    elif document_type == 'OPENING_MINUTES':
        _fill_opening_minutes(doc, fields)
    else:
        raise ValidationError('این منبع DOCX برای نوع سند درخواستی تعریف نشده است.')
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), 'docx'
