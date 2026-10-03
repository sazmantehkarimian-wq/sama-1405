"""Single official output engine for filtered canonical querysets."""
from io import BytesIO
from pathlib import Path
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.utils import get_column_letter
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
import arabic_reshaper
from bidi.algorithm import get_display

HEADERS=['سازمان فرهنگی هنری شهرداری تهران','مدیریت اقتصادی و املاک','اداره املاک و مستغلات']
FIELD_MAP={
 'code':('کد فضا',lambda s:s.code),'name':('نام فضا / مرکز',lambda s:s.name or '—'),
 'status':('وضعیت',lambda s:s.get_status_display()),'region':('منطقه',lambda s:s.region.name if s.region else '—'),
 'current_usage':('کاربری',lambda s:s.current_usage or '—'),'area':('مساحت (مترمربع)',lambda s:s.area if s.area is not None else '—'),
}
MOTHER_FIELD_MAP={
 'identifier':('شناسه ملک',lambda p:p.identifier),'name':('نام ملک / مرکز',lambda p:p.name or '—'),
 'region':('منطقه',lambda p:p.region.name if p.region else '—'),'address':('نشانی',lambda p:p.address or '—'),
 'primary_usage':('کاربری اصلی',lambda p:p.primary_usage or '—'),'usage_group':('گروه کاربری',lambda p:p.usage_group or '—'),
 'area':('مساحت (مترمربع)',lambda p:p.area if p.area is not None else '—'),
 'space_count':('تعداد فضاهای مرتبط',lambda p:getattr(p,'space_count',p.space_links.count())),
}
DEFAULT_FIELDS=tuple(FIELD_MAP)
BASE=Path(__file__).resolve().parents[1]
LOGO=BASE/'design_system/static/design_system/img/organization-logo.png'
FONT=BASE/'design_system/static/design_system/fonts/Vazirmatn-Regular.ttf'

def _style_xlsx(ws,header,total):
 ws.freeze_panes=f'A{header+1}';ws.auto_filter.ref=f'A{header}:{ws.cell(max(header,ws.max_row),total).coordinate}'
 ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.print_title_rows=f'{header}:{header}';ws.print_area=f'A1:{get_column_letter(total)}{max(header,ws.max_row)}';ws.oddFooter.center.text='&N از &P صفحه'
 side=Side(style='thin',color='D8D3CF');header_fill=PatternFill('solid',fgColor='E7F1F4')
 for row in ws.iter_rows(min_row=header):
  for cell in row:
   cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True,readingOrder=2);cell.border=Border(left=side,right=side,top=side,bottom=side);cell.font=Font(name='Vazirmatn',size=10)
   if isinstance(cell.value,(int,float,Decimal)):cell.number_format='#,##0.###'
 for cell in ws[header]:cell.font=Font(name='Vazirmatn',bold=True,size=10,color='252221');cell.fill=header_fill
 for row_number in range(1,header):
  ws.row_dimensions[row_number].height=24
  for cell in ws[row_number]:cell.alignment=Alignment(horizontal='center',vertical='center',readingOrder=2);cell.font=Font(name='Vazirmatn',bold=True,size=11)
 ws.row_dimensions[header].height=28
 for column in ws.columns:ws.column_dimensions[get_column_letter(column[0].column)].width=min(42,max(14,max(len(str(cell.value or '')) for cell in column)+2))

def columns(selected=None,field_map=None):
 field_map=field_map or FIELD_MAP;selected=[key for key in (selected or tuple(field_map)) if key in field_map]
 return selected or list(field_map)
def layout_columns(layout=None,selected=None,blank_columns=(),field_map=None):
 """Return one authoritative ordered output layout (first item is visual right)."""
 field_map=field_map or FIELD_MAP;result=[]
 for token in layout or ():
  kind,separator,value=str(token).partition(':')
  if separator and kind=='field' and value in field_map and ('field',value) not in result:result.append(('field',value))
  elif separator and kind=='blank' and value.strip():result.append(('blank',value.strip()[:120]))
 if not result:result=[('field',key) for key in columns(selected,field_map)]+[('blank',str(label).strip()[:120]) for label in blank_columns if str(label).strip()]
 return result[:24]
def _blank_row_count(value):
 try:return max(0,min(50,int(value or 0)))
 except (TypeError,ValueError):return 0
def output_table(spaces,layout=None,selected=None,blank_columns=(),field_map=None,blank_rows=0):
 field_map=field_map or FIELD_MAP;configured=layout_columns(layout,selected,blank_columns,field_map)
 headers=[field_map[value][0] if kind=='field' else value for kind,value in configured]
 data=[[field_map[value][1](space) if kind=='field' else '' for kind,value in configured] for space in spaces]
 data.extend([['' for _ in headers] for _ in range(_blank_row_count(blank_rows))])
 return headers,data
def rows(spaces,selected=None):
 keys=columns(selected)
 for s in spaces: yield [FIELD_MAP[key][1](s) for key in keys]
def excel(spaces,blank_columns=(),selected=None,layout=None,orientation='landscape',field_map=None,title='فضاها',blank_rows=0):
 labels,data=output_table(spaces,layout,selected,blank_columns,field_map,blank_rows=blank_rows); wb=Workbook(); ws=wb.active; ws.title=title[:31]; ws.sheet_view.rightToLeft=True
 total=len(labels)
 for line in HEADERS:
  ws.append(['',line] if total>1 else [line])
  if total>2:ws.merge_cells(start_row=ws.max_row,start_column=2,end_row=ws.max_row,end_column=total)
 if LOGO.exists():
  logo=ExcelImage(str(LOGO));logo.width=64;logo.height=64;ws.add_image(logo,'A1')
 ws.append(labels); header=ws.max_row
 for row in data:ws.append(row)
 ws.page_setup.orientation=orientation
 _style_xlsx(ws,header,total)
 out=BytesIO(); wb.save(out); return out.getvalue()

def tabular_excel(title,labels,data):
 """Official Excel 2019-compatible output for typed operational domain reports."""
 wb=Workbook();ws=wb.active;ws.title=str(title)[:31];ws.sheet_view.rightToLeft=True
 total=max(1,len(labels))
 for line in HEADERS:
  ws.append(['',line] if total>1 else [line])
  if total>2:ws.merge_cells(start_row=ws.max_row,start_column=2,end_row=ws.max_row,end_column=total)
 if LOGO.exists():
  logo=ExcelImage(str(LOGO));logo.width=64;logo.height=64;ws.add_image(logo,'A1')
 ws.append(list(labels));header=ws.max_row
 for values in data:ws.append(list(values))
 _style_xlsx(ws,header,total)
 out=BytesIO();wb.save(out);return out.getvalue()
def _rtl(paragraph):
 paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pPr=paragraph._p.get_or_add_pPr();bidi=OxmlElement('w:bidi');bidi.set(qn('w:val'),'1');pPr.append(bidi)
 for run in paragraph.runs:
  run.font.name='Vazirmatn';run.font.size=Pt(10);run._element.rPr.rFonts.set(qn('w:cs'),'Vazirmatn')
def _rtl_table(table):
 tblPr=table._tbl.tblPr;bidi=OxmlElement('w:bidiVisual');bidi.set(qn('w:val'),'1');tblPr.append(bidi)
def docx(spaces,blank_columns=(),selected=None,layout=None,orientation='landscape',field_map=None,blank_rows=0):
 labels,data=output_table(spaces,layout,selected,blank_columns,field_map,blank_rows=blank_rows); doc=Document();sec=doc.sections[0]
 if orientation=='landscape':sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width,sec.page_height=sec.page_height,sec.page_width
 sec.top_margin=Mm(40);sec.header_distance=Mm(5)
 header=sec.header.paragraphs[0]
 if LOGO.exists(): header.add_run().add_picture(str(LOGO),width=Mm(18))
 header.add_run('\n'+'\n'.join(HEADERS));_rtl(header)
 table=doc.add_table(rows=1,cols=len(labels));table.style='Table Grid'
 _rtl_table(table)
 for i,h in enumerate(labels):table.rows[0].cells[i].text=h;_rtl(table.rows[0].cells[i].paragraphs[0])
 table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for row in data:
  cells=table.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v);_rtl(cells[i].paragraphs[0])
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
 footer.add_run('صفحه ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 out=BytesIO();doc.save(out);return out.getvalue()
def _fa(value):return get_display(arabic_reshaper.reshape(str(value).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))))
def pdf(spaces,selected=None,blank_columns=(),layout=None,orientation='landscape',field_map=None,blank_rows=0):
 labels,body=output_table(spaces,layout,selected,blank_columns,field_map,blank_rows=blank_rows);out=BytesIO();pdfmetrics.registerFont(TTFont('Vazirmatn',str(FONT)));page_size=landscape(A4) if orientation=='landscape' else A4
 doc=SimpleDocTemplate(out,pagesize=page_size,rightMargin=12*mm,leftMargin=12*mm,topMargin=10*mm,bottomMargin=12*mm,title='گزارش رسمی املاک')
 style=ParagraphStyle('fa',fontName='Vazirmatn',fontSize=9,leading=14,alignment=TA_CENTER)
 story=[]
 if LOGO.exists():story.append(Image(str(LOGO),width=18*mm,height=18*mm))
 for line in HEADERS:story.append(Paragraph(_fa(line),style))
 story.append(Spacer(1,5*mm))
 data=[[Paragraph(_fa(value),style) for value in labels]]+[[Paragraph(_fa(value),style) if value!='' else '' for value in row] for row in body]
 data=[list(reversed(row)) for row in data]
 available=page_size[0]-24*mm;table=Table(data,repeatRows=1,hAlign='CENTER',splitByRow=1,colWidths=[available/max(1,len(labels))]*len(labels))
 table.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'Vazirmatn'),('FONTSIZE',(0,0),(-1,-1),8),('ALIGN',(0,0),(-1,-1),'CENTER'),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('GRID',(0,0),(-1,-1),.5,colors.grey),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#f1edf2')),('LEADING',(0,0),(-1,-1),12)]));story.append(table)
 def numbered_page(canvas,document):
  canvas.saveState();canvas.setFont('Vazirmatn',8);canvas.drawCentredString(page_size[0]/2,6*mm,_fa(f'صفحه {document.page}'));canvas.restoreState()
 doc.build(story,onFirstPage=numbered_page,onLaterPages=numbered_page);return out.getvalue()