import io
import pytest
from openpyxl import load_workbook
from domains.properties.models import CommercialSpace,Region
from reporting.engine import excel,output_table
pytestmark=pytest.mark.django_db

def test_layout_blank_columns_multi_sort_and_rtl():
 r1=Region.objects.create(code='2',name='دو');r2=Region.objects.create(code='1',name='یک')
 CommercialSpace.objects.create(code='2',name='ب',status='ACTIVE',region=r1,source_row=2,source_classification='AUTHORITY');CommercialSpace.objects.create(code='1',name='الف',status='ACTIVE',region=r2,source_row=1,source_classification='AUTHORITY')
 qs=CommercialSpace.objects.select_related('region').order_by('region__name','code');layout=['field:code','blank:توضیحات کارشناس','field:region','blank:امضاء']
 labels,rows=output_table(qs,layout)
 assert labels==['کد فضا','توضیحات کارشناس','منطقه','امضاء'] and [row[0] for row in rows]==['2','1'] and all(row[1]==row[3]=='' for row in rows)
 wb=load_workbook(io.BytesIO(excel(qs,layout=layout)));ws=wb.active
 assert ws.sheet_view.rightToLeft is True and [cell.value for cell in ws[4]]==labels and ws.freeze_panes=='A5' and ws.auto_filter.ref


def test_blank_output_rows_are_output_only():
 from reporting.engine import output_table
 class Space:
  code='1';name='نمونه';area=None;current_usage='';region=None
  def get_status_display(self): return 'فعال'
 labels,rows=output_table([Space()],selected=['code','name'],blank_rows=3)
 assert labels==['کد فضا','نام فضا / مرکز']
 assert rows==[['1','نمونه'],['',''],['',''],['','']]
