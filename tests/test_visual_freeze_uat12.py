import io
import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook

from domains.properties.models import Region, Center, MotherProperty, CommercialSpace
from reporting.engine import output_table, FIELD_MAP

pytestmark = pytest.mark.django_db


def login(client):
    user = get_user_model().objects.create_user('visual-user')
    client.force_login(user)
    return user


def test_regions_workspace_and_special_centers_are_live_aggregations(client):
    login(client)
    r1 = Region.objects.create(code='1', name='منطقه 1')
    MotherProperty.objects.create(identifier='P-1', name='ملک یک', region=r1, source_row=1)
    regular = Center.objects.create(name='مرکز عادی', region=r1, is_special=False)
    special = Center.objects.create(name='مرکز خاص نمونه', region=r1, is_special=True)
    CommercialSpace.objects.create(code='2', name='فضای عادی', status='ACTIVE', region=r1, center=regular, source_row=1, source_classification='authority')
    CommercialSpace.objects.create(code='3', name='فضای خاص', status='OUT_OF_CYCLE', region=r1, center=special, source_row=2, source_classification='authority')
    response = client.get('/regions/1/')
    body = response.content.decode()
    assert response.status_code == 200 and 'منطقه 1 — نمای مدیریتی' in body
    assert 'ملک یک' in body and 'فضای عادی' in body and 'فضای خاص' in body
    special_response = client.get(f'/regions/special/{special.pk}/')
    special_body = special_response.content.decode()
    assert special_response.status_code == 200 and 'مرکز خاص نمونه — پرونده مرکز خاص' in special_body
    assert 'فضای خاص' in special_body and 'فضای عادی' not in special_body


def test_central_report_builder_exposes_mother_properties_and_preserves_context(client):
    login(client)
    region = Region.objects.create(code='1', name='منطقه 1')
    response = client.get('/reports/', {'domain': 'mother_properties', 'identifier': 'P-22', 'region_id': region.pk, 'blank_rows': 3, 'orientation': 'portrait'})
    body = response.content.decode()
    assert response.status_code == 200
    assert 'املاک مادر' in body and 'value="P-22"' in body
    assert f'value="{region.pk}" selected' in body
    assert 'value="3" selected' in body and 'value="portrait" selected' in body


def test_blank_output_rows_are_output_only_and_xlsx_orientation_is_respected(client):
    login(client)
    region = Region.objects.create(code='1', name='منطقه 1')
    CommercialSpace.objects.create(code='1', name='فضای یک', status='ACTIVE', region=region, source_row=1, source_classification='authority')
    labels, rows = output_table(CommercialSpace.objects.all(), selected=['code', 'name'], field_map=FIELD_MAP, blank_rows=3)
    assert labels == ['کد فضا', 'نام فضا / مرکز']
    assert len(rows) == 4 and rows[-1] == ['', ''] and CommercialSpace.objects.count() == 1
    response = client.get('/reports/spaces.xlsx', {'field': ['code', 'name'], 'blank_rows': '3', 'orientation': 'portrait'})
    ws = load_workbook(io.BytesIO(response.content)).active
    assert ws.sheet_view.rightToLeft is True and ws.page_setup.orientation == 'portrait'
    assert ws.max_row == 8  # three official header rows + table header + one data row + three blank output rows
    assert CommercialSpace.objects.count() == 1


def test_shared_picker_closes_after_selection_and_supports_escape_and_click_outside():
    script = open('design_system/static/design_system/js/app.js', encoding='utf-8').read()
    assert "picker.classList.remove('open')" in script
    assert "event.key==='Escape'" in script
    assert "pointerdown" in script
    assert "choose(item)" in script


def test_every_authenticated_page_gets_central_print_action(client):
    login(client)
    Region.objects.create(code='1', name='منطقه 1')
    for path in ('/', '/mother-properties/', '/spaces/', '/reports/', '/regions/'):
        response = client.get(path)
        assert response.status_code == 200
        assert 'data-page-print' in response.content.decode()
