import io, zipfile
import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from pypdf import PdfReader
from domains.contracts.models import Beneficiary, BeneficiaryAssignment, Contract
from domains.operations.models import Appraisal, UtilityRecord
from domains.properties.models import CommercialSpace, MotherProperty, Region
from reporting.dossier_sections import section_report

@pytest.fixture
def user(db):return get_user_model().objects.create_user('phase1',password='safe-pass')

@pytest.mark.django_db
def test_dossier_section_dataset_and_four_formats_are_isolated(client,user):
 space=CommercialSpace.objects.create(code='347',name='فضای ۳۴۷',status='ACTIVE',source_row=1,source_classification='authority')
 Contract.objects.create(space=space,number='C-1',is_historical=True,source_row=2)
 for index in range(3):Appraisal.objects.create(space=space,year='1404',amount_rial=0 if index==0 else None,appraiser='' if index==1 else 'کارشناس',appraisal_date='' if index==2 else f'1404/01/0{index+1}',status='قابل استفاده با نقص اطلاعات',source_row=index+3)
 client.force_login(user);report=section_report(space,'appraisals')
 assert len(report.rows)==3 and 'سوابق کارشناسی فضای تجاری 347'==report.title
 for extension in ('xlsx','pdf','docx','print'):
  response=client.get(f'/spaces/347/sections/appraisals.{extension}')
  assert response.status_code==200
 xlsx=load_workbook(io.BytesIO(client.get('/spaces/347/sections/appraisals.xlsx?blank_rows=2').content)).active
 assert xlsx.sheet_view.rightToLeft and xlsx.max_row==9
 pdf=client.get('/spaces/347/sections/appraisals.pdf').content
 assert PdfReader(io.BytesIO(pdf)).metadata.title=='سوابق کارشناسی فضای تجاری 347'
 with zipfile.ZipFile(io.BytesIO(client.get('/spaces/347/sections/appraisals.docx').content)) as archive:
  xml=''.join(archive.read(name).decode(errors='ignore') for name in archive.namelist() if name.endswith('.xml'))
  assert 'w:bidiVisual' in xml and 'سوابق کارشناسی فضای تجاری 347' in xml and 'SMK' not in xml and 'سما' not in xml
 html=client.get('/spaces/347/sections/appraisals.print').content.decode()
 assert 'سوابق کارشناسی فضای تجاری 347' in html and 'بهره‌بردار' not in html and 'قراردادها' not in html and 'SMK' not in html and 'سما' not in html and '<button' not in html

@pytest.mark.django_db
def test_dossier_and_specialist_counts_match_for_known_examples(client,user):
 s347=CommercialSpace.objects.create(code='347',name='نمونه ۳۴۷',status='ACTIVE',source_row=1,source_classification='authority')
 s497=CommercialSpace.objects.create(code='497',name='نمونه ۴۹۷',status='ACTIVE',source_row=2,source_classification='authority')
 for space,count in ((s347,3),(s497,2)):
  Contract.objects.create(space=space,number='',is_historical=True,source_row=1)
  for index in range(count):Appraisal.objects.create(space=space,amount_rial=0 if index==0 else None,appraisal_date='' if index==0 else '1404/01/01',appraiser='' if index==0 else 'کارشناس',status='قابل استفاده با نقص اطلاعات' if index==0 else '',source_row=index+1)
 client.force_login(user)
 for code,count in (('347',3),('497',2)):
  assert client.get('/records/appraisals/',{'space_code':code}).context['page'].paginator.count==count
  assert len(section_report(CommercialSpace.objects.get(code=code),'appraisals').rows)==count
  assert client.get('/records/contracts/',{'space_code':code}).context['page'].paginator.count==1
  assert len(section_report(CommercialSpace.objects.get(code=code),'contracts').rows)==1

@pytest.mark.django_db
def test_appraisal_zero_null_quality_and_filtered_exports_match(client,user):
 region=Region.objects.create(code='16',name='منطقه ۱۶');space=CommercialSpace.objects.create(code='37',name='فضا',status='ACTIVE',region=region,source_row=1,source_classification='authority')
 Appraisal.objects.create(space=space,amount_rial=0,appraisal_date='',appraiser='',status='قابل استفاده با نقص اطلاعات',source_row=1)
 Appraisal.objects.create(space=space,amount_rial=None,appraisal_date='1404/01/01',appraiser='الف',status='',source_row=2)
 Appraisal.objects.create(space=space,amount_rial=100,appraisal_date='1404/02/01',appraiser='ب',status='',source_row=3)
 client.force_login(user)
 zero=client.get('/records/appraisals/',{'space_code':'37','quality':'zero'});unknown=client.get('/records/appraisals/',{'space_code':'37','quality':'unknown'})
 assert zero.context['page'].paginator.count==1 and unknown.context['page'].paginator.count==1
 response=client.get('/records/appraisals/export.xlsx',{'space_code':'37','quality':'positive','sort':'amount','column':['space.code','amount_rial']})
 sheet=load_workbook(io.BytesIO(response.content)).active
 assert sheet.max_row==5 and [sheet.cell(4,c).value for c in (1,2)]==['کد فضا','مبلغ (ریال)'] and sheet.cell(5,2).value==100

@pytest.mark.django_db
def test_region_property_counts_are_explicit_and_reconcile(client,user):
 mapped=Region.objects.create(code='1',name='منطقه ۱');MotherProperty.objects.create(identifier='M1',name='مرتبط',region=mapped,source_row=1);MotherProperty.objects.create(identifier='M2',name='نامشخص',source_row=2)
 client.force_login(user);response=client.get('/regions/')
 assert response.context['total_property_count']==2 and response.context['mapped_property_count']==1 and response.context['unmapped_property_count']==1
 assert response.context['mapped_property_count']+response.context['unmapped_property_count']==response.context['total_property_count']
