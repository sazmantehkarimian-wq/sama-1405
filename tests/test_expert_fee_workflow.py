from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook

from domains.identity.models import AuditEvent
from domains.operations.models import (
    Appraiser, Appraisal, AppraisalFee, ExpertFeeBatchItem, ExpertFeePaymentBatch,
)
from domains.properties.models import CommercialSpace


def fee_context():
    user=get_user_model().objects.create_user("fee-user",password="A-very-safe-password")
    space=CommercialSpace.objects.create(code="8301",name="فضای حق‌الزحمه",status="ACTIVE")
    appraiser=Appraiser.objects.create(first_name="ندا",last_name="کارشناس",created_by=user)
    appraisal=Appraisal.objects.create(
        space=space,appraiser_ref=appraiser,appraiser=appraiser.full_name,
        appraisal_date="1405/07/01",amount_rial=5000000,status="تکمیل‌شده",
        created_by=user,
    )
    return user,space,appraiser,appraisal


@pytest.mark.django_db
def test_fee_is_manual_independent_from_appraisal_amount_and_audited(client):
    user,space,appraiser,appraisal=fee_context()
    client.force_login(user)
    response=client.post(f"/appraisals/{appraisal.pk}/fees/new/",{
        "amount_rial":"1250000","follow_up_date":"1405/07/20","notes":"مبلغ توافق‌شده",
    })
    assert response.status_code==302
    fee=AppraisalFee.objects.get(appraisal=appraisal)
    assert fee.amount_rial==1250000
    assert fee.amount_rial!=appraisal.amount_rial
    assert fee.status==AppraisalFee.Status.FEE_ENTERED
    assert fee.expert==appraiser and fee.space==space
    assert AuditEvent.objects.filter(action="APPRAISAL_FEE_CREATE",entity_id=str(fee.pk)).exists()


@pytest.mark.django_db
def test_fee_sent_to_finance_requires_letter_and_full_payment_matches_fee(client):
    user,_,_,appraisal=fee_context()
    client.force_login(user)
    client.post(f"/appraisals/{appraisal.pk}/fees/new/",{"amount_rial":"1200000"})
    fee=AppraisalFee.objects.get()

    assert client.post(f"/fees/{fee.pk}/transition/",{"status":"READY_TO_SEND"}).status_code==302
    fee.refresh_from_db()
    assert fee.status==AppraisalFee.Status.READY_TO_SEND

    rejected=client.post(f"/fees/{fee.pk}/transition/",{
        "status":"SENT_TO_FINANCE","sent_to_finance_date":"1405/07/10",
    })
    assert rejected.status_code==200
    fee.refresh_from_db()
    assert fee.status==AppraisalFee.Status.READY_TO_SEND

    sent=client.post(f"/fees/{fee.pk}/transition/",{
        "status":"SENT_TO_FINANCE","sent_to_finance_date":"1405/07/10",
        "letter_number":"FIN-10","letter_date":"1405/07/10",
    })
    assert sent.status_code==302
    assert client.post(f"/fees/{fee.pk}/transition/",{"status":"IN_PROGRESS"}).status_code==302

    mismatch=client.post(f"/fees/{fee.pk}/transition/",{
        "status":"PAID","payment_date":"1405/08/01","paid_amount_rial":"1100000","payment_reference":"PAY-REF",
    })
    assert mismatch.status_code==200
    fee.refresh_from_db()
    assert fee.status==AppraisalFee.Status.IN_PROGRESS

    paid=client.post(f"/fees/{fee.pk}/transition/",{
        "status":"PAID","payment_date":"1405/08/01","paid_amount_rial":"1200000","payment_reference":"PAY-REF",
    })
    assert paid.status_code==302
    fee.refresh_from_db()
    assert fee.status==AppraisalFee.Status.PAID
    assert fee.paid_amount_rial==fee.amount_rial
    assert fee.pending_amount_rial==0

    blocked=client.post(f"/fees/{fee.pk}/amount/",{"amount_rial":"1300000","reason":"اصلاح"})
    assert blocked.status_code==200
    fee.refresh_from_db()
    assert fee.amount_rial==1200000

    assert client.post(f"/fees/{fee.pk}/transition/",{"status":"CLOSED"}).status_code==302
    fee.refresh_from_db()
    assert fee.status==AppraisalFee.Status.CLOSED


@pytest.mark.django_db
def test_fee_amount_correction_before_payment_requires_reason(client):
    user,_,_,appraisal=fee_context()
    client.force_login(user)
    client.post(f"/appraisals/{appraisal.pk}/fees/new/",{"amount_rial":"900000"})
    fee=AppraisalFee.objects.get()

    no_reason=client.post(f"/fees/{fee.pk}/amount/",{"amount_rial":"950000","reason":""})
    assert no_reason.status_code==200
    fee.refresh_from_db()
    assert fee.amount_rial==900000

    changed=client.post(f"/fees/{fee.pk}/amount/",{"amount_rial":"950000","reason":"اصلاح مبلغ مصوب"})
    assert changed.status_code==302
    fee.refresh_from_db()
    assert fee.amount_rial==950000
    audit=AuditEvent.objects.get(action="APPRAISAL_FEE_AMOUNT_UPDATE",entity_id=str(fee.pk))
    assert audit.reason=="اصلاح مبلغ مصوب"


@pytest.mark.django_db
def test_fee_batch_sends_only_ready_records_and_enforces_single_active_batch(client):
    user,space,appraiser,first_appraisal=fee_context()
    second_appraisal=Appraisal.objects.create(
        space=space,appraiser_ref=appraiser,appraiser=appraiser.full_name,
        appraisal_date="1405/08/01",amount_rial=6000000,created_by=user,
    )
    client.force_login(user)
    client.post(f"/appraisals/{first_appraisal.pk}/fees/new/",{"amount_rial":"1000000"})
    client.post(f"/appraisals/{second_appraisal.pk}/fees/new/",{"amount_rial":"1500000"})
    fees=list(AppraisalFee.objects.order_by("pk"))
    for fee in fees:
        assert client.post(f"/fees/{fee.pk}/transition/",{"status":"READY_TO_SEND"}).status_code==302

    response=client.post("/fees/batches/new/",{
        "fees":[str(fee.pk) for fee in fees],
        "sent_date":"1405/08/10","letter_number":"BATCH-LETTER","letter_date":"1405/08/10",
        "notes":"ارسال گروهی",
    })
    assert response.status_code==302
    batch=ExpertFeePaymentBatch.objects.get()
    assert batch.sama_code.startswith("PAY-1405-")
    assert batch.items.filter(active=True).count()==2
    assert batch.total_amount_rial==2500000
    assert ExpertFeeBatchItem.objects.filter(active=True).count()==2
    assert all(f.status==AppraisalFee.Status.SENT_TO_FINANCE for f in AppraisalFee.objects.all())
    assert AuditEvent.objects.filter(action="APPRAISAL_FEE_BATCH_CREATE",entity_id=str(batch.pk)).exists()


@pytest.mark.django_db
def test_fee_dashboard_filters_and_excel_export(client):
    user,space,appraiser,appraisal=fee_context()
    client.force_login(user)
    client.post(f"/appraisals/{appraisal.pk}/fees/new/",{"amount_rial":"700000"})
    fee=AppraisalFee.objects.get()
    client.post(f"/fees/{fee.pk}/transition/",{"status":"READY_TO_SEND"})

    response=client.get("/fees/",{"status":"READY_TO_SEND"})
    assert response.status_code==200
    assert response.context["fee_count"]==1
    assert fee.sama_code in response.content.decode()
    assert appraiser.full_name in response.content.decode()

    xlsx=client.get("/fees/export.xlsx",{"q":space.code})
    assert xlsx.status_code==200
    workbook=load_workbook(BytesIO(xlsx.content),data_only=True)
    sheet=workbook.active
    values=[cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert fee.sama_code in values
    assert appraiser.full_name in values
    assert sheet.sheet_view.rightToLeft is True

    pdf=client.get("/fees/export.pdf",{"q":space.code})
    assert pdf.status_code==200
    assert pdf["Content-Type"]=="application/pdf"
    assert pdf.content.startswith(b"%PDF")


@pytest.mark.django_db
def test_fee_creation_requires_registered_appraiser(client):
    user=get_user_model().objects.create_user("fee-no-expert",password="A-very-safe-password")
    client.force_login(user)
    space=CommercialSpace.objects.create(code="8302",name="فضا",status="ACTIVE")
    appraisal=Appraisal.objects.create(space=space,appraiser="نام آزاد",appraisal_date="1405/07/01",amount_rial=100)
    response=client.post(f"/appraisals/{appraisal.pk}/fees/new/",{"amount_rial":"1000"})
    assert response.status_code==200
    assert AppraisalFee.objects.count()==0
    assert "کارشناس ثبت‌شده" in response.content.decode()
