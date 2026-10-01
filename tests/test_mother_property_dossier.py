import io
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile

from domains.documents.models import Document
from domains.identity.models import AuditEvent
from domains.properties.models import (
    MotherProperty, MotherPropertyCorrespondence, MotherPropertyNote,
    MotherPropertyOwnership, MotherPropertyOwnershipDocument,
    MotherPropertyUsageHistory, PropertyReferenceValue,
)


def add_reference(category, value):
    return PropertyReferenceValue.objects.create(category=category, value=value, active=True)


@pytest.fixture
def property_refs(db):
    add_reference("STATUS", "فعال")
    add_reference("CENTER_TYPE", "فرهنگسرا")
    add_reference("USAGE", "فرهنگی")
    add_reference("USAGE_GROUP", "فرهنگی و هنری")
    add_reference("ORG_UNIT", "منطقه ۵")
    add_reference("USAGE_STATUS", "در اختیار واحد سازمانی")
    add_reference("USAGE_STATUS", "دارای بهره‌بردار")
    add_reference("OWNERSHIP_STATUS", "موجود است")


@pytest.fixture
def property_user(db):
    return get_user_model().objects.create_user("property-user", password="A-very-safe-password")


@pytest.mark.django_db
def test_mother_property_create_uses_controlled_reference_data_and_stays_independent(client, property_user, property_refs):
    client.force_login(property_user)
    response=client.post("/properties/new/",{
        "identifier":"P-0100","name":"ملک مستقل","current_status":"فعال",
        "center_type":"فرهنگسرا","primary_usage":"فرهنگی","usage_group":"فرهنگی و هنری",
        "holder_unit":"منطقه ۵","land_area":"1000.50","area":"450.25",
        "ownership_document_status":"موجود است","owner_name":"سازمان","owner_type":"ORGANIZATIONAL",
        "has_utilities":"YES","electricity_presence":"YES","water_presence":"YES","gas_presence":"NO",
    })
    assert response.status_code==302
    item=MotherProperty.objects.get(identifier="P-0100")
    assert response.url.endswith(f"/properties/{item.pk}/")
    assert item.created_by==property_user and item.updated_by==property_user
    assert item.land_area is not None and item.area is not None
    assert item.completeness_status=="کامل"
    assert not hasattr(item, "space_links")
    assert AuditEvent.objects.filter(action="MOTHER_PROPERTY_CREATE",entity_id="P-0100").exists()

    dossier=client.get(response.url)
    assert dossier.status_code==200
    body=dossier.content.decode()
    assert "این پرونده مستقل از کدهای فضای تجاری است" in body
    assert 'id="appraisals"' not in body
    assert 'id="auctions"' not in body


@pytest.mark.django_db
def test_mother_property_rejects_free_text_for_reference_fields(client, property_user, property_refs):
    client.force_login(property_user)
    response=client.post("/properties/new/",{
        "identifier":"P-0101","name":"ملک","current_status":"وضعیت دلخواه",
    })
    assert response.status_code==200
    assert MotherProperty.objects.count()==0
    assert "Select a valid choice" in response.content.decode() or "انتخاب" in response.content.decode()


@pytest.mark.django_db
def test_mother_property_usage_history_closes_previous_only_with_reason(client, property_user, property_refs):
    client.force_login(property_user)
    item=MotherProperty.objects.create(identifier="P-0102",name="ملک",current_status="فعال",created_by=property_user)

    first=client.post(f"/properties/{item.pk}/usage/change/",{
        "usage_status":"در اختیار واحد سازمانی","holder_type":"REGION","holder_unit":"منطقه ۵",
        "start_date":"1405/07/01","basis":"تحویل سازمانی",
    })
    assert first.status_code==302
    current=MotherPropertyUsageHistory.objects.get(property=item)
    assert current.end_date==""

    rejected=client.post(f"/properties/{item.pk}/usage/change/",{
        "usage_status":"دارای بهره‌بردار","beneficiary_name":"شرکت نمونه",
        "start_date":"1405/08/01","basis":"تحویل بهره‌برداری",
    })
    assert rejected.status_code==200
    assert MotherPropertyUsageHistory.objects.filter(property=item).count()==1

    accepted=client.post(f"/properties/{item.pk}/usage/change/",{
        "usage_status":"دارای بهره‌بردار","beneficiary_name":"شرکت نمونه","beneficiary_type":"LEGAL",
        "start_date":"1405/08/01","basis":"تحویل بهره‌برداری","termination_reason":"تغییر بهره‌برداری",
    })
    assert accepted.status_code==302
    history=list(MotherPropertyUsageHistory.objects.filter(property=item).order_by("start_date"))
    assert len(history)==2
    assert history[0].end_date=="1405/07/30"
    assert history[0].termination_reason=="تغییر بهره‌برداری"
    assert history[1].end_date==""
    assert AuditEvent.objects.filter(action="MOTHER_PROPERTY_USAGE_CHANGE",entity_id=item.identifier).count()==2


@pytest.mark.django_db
def test_mother_property_ownership_supports_multiple_owners_and_separate_document_area(client, property_user):
    client.force_login(property_user)
    item=MotherProperty.objects.create(identifier="P-0103",name="ملک",created_by=property_user)

    for owner,share in (("مالک الف","60"),("مالک ب","40")):
        response=client.post(f"/properties/{item.pk}/ownership/new/",{
            "owner_name":owner,"owner_type":"NATURAL","share_percent":share,"start_date":"1405/01/01",
        })
        assert response.status_code==302
    assert MotherPropertyOwnership.objects.filter(property=item).count()==2

    response=client.post(f"/properties/{item.pk}/ownership-documents/new/",{
        "document_type":"سند رسمی","document_number":"DOC-1","document_date":"1405/02/01",
        "documented_area":"999.75","documented_owner_name":"مالک الف","owner_type":"NATURAL","status":"ACTIVE",
    })
    assert response.status_code==302
    document=MotherPropertyOwnershipDocument.objects.get(property=item)
    assert str(document.documented_area)=="999.75"
    item.refresh_from_db()
    assert item.area is None and item.land_area is None


@pytest.mark.django_db
def test_mother_property_correspondence_followup_requires_owner_and_due_date(client, property_user):
    client.force_login(property_user)
    item=MotherProperty.objects.create(identifier="P-0104",name="ملک",created_by=property_user)

    invalid=client.post(f"/properties/{item.pk}/correspondence/new/",{
        "document_type":"INCOMING","subject":"نامه پیگیری","needs_follow_up":"on","follow_up_status":"OPEN",
    })
    assert invalid.status_code==200
    assert MotherPropertyCorrespondence.objects.count()==0

    valid=client.post(f"/properties/{item.pk}/correspondence/new/",{
        "document_type":"INCOMING","subject":"نامه پیگیری","needs_follow_up":"on",
        "responsible":property_user.pk,"due_date":"1405/09/01","follow_up_status":"OPEN",
    })
    assert valid.status_code==302
    record=MotherPropertyCorrespondence.objects.get(property=item)
    assert record.responsible==property_user and record.due_date=="1405/09/01"


@pytest.mark.django_db
def test_mother_property_notes_are_append_only_and_documents_archive_without_delete(client, property_user, tmp_path, settings):
    settings.MEDIA_ROOT=tmp_path
    client.force_login(property_user)
    item=MotherProperty.objects.create(identifier="P-0105",name="ملک",created_by=property_user)

    assert client.post(f"/properties/{item.pk}/notes/new/",{"subject":"اول","text":"یادداشت اول","active":"on"}).status_code==302
    assert client.post(f"/properties/{item.pk}/notes/new/",{"subject":"دوم","text":"یادداشت دوم","active":"on"}).status_code==302
    assert MotherPropertyNote.objects.filter(property=item).count()==2

    upload=client.post(f"/properties/{item.pk}/documents/",{
        "title":"نامه ملک","document_type":"نامه","reference":"MP-1","document_date":"1405/07/01",
        "file":SimpleUploadedFile("property.pdf",b"%PDF-1.4\nSAMA",content_type="application/pdf"),
    })
    assert upload.status_code==302
    document=Document.objects.get(entity_type="MotherProperty",entity_id=item.identifier)
    archive=client.post(f"/documents/{document.pk}/archive/",{"reason":"نسخه قدیمی"})
    assert archive.status_code==302
    assert archive.url.endswith(f"/properties/{item.pk}/")
    document.refresh_from_db()
    assert document.archived_at is not None
    assert Document.objects.filter(pk=document.pk).exists()
    dossier=client.get(f"/properties/{item.pk}/").content.decode()
    assert "بایگانی‌شده" in dossier and "MP-1" in dossier



@pytest.mark.django_db
def test_mother_property_official_exports_are_rtl_and_audited(client, property_user, property_refs):
    client.force_login(property_user)
    item=MotherProperty.objects.create(
        identifier="P-0106",name="ملک گزارش",current_status="فعال",
        primary_usage="فرهنگی",area=450,ownership_document_status="موجود است",
        electricity_presence="YES",water_presence="YES",gas_presence="NO",
        created_by=property_user,updated_by=property_user,
    )

    xlsx=client.get(f"/properties/{item.pk}/export.xlsx")
    assert xlsx.status_code==200
    workbook=load_workbook(io.BytesIO(xlsx.content),data_only=True)
    assert len(workbook.sheetnames)>=5
    assert all(workbook[name].sheet_view.rightToLeft for name in workbook.sheetnames)
    values=[cell.value for name in workbook.sheetnames for row in workbook[name].iter_rows() for cell in row if cell.value is not None]
    assert "سازمان فرهنگی هنری شهرداری تهران" in values
    assert "P-0106" in [str(v) for v in values]
    assert "سامانه مدیریت قراردادها" not in values

    pdf=client.get(f"/properties/{item.pk}/export.pdf")
    assert pdf.status_code==200 and pdf.content.startswith(b"%PDF")
    reader=PdfReader(io.BytesIO(pdf.content))
    assert reader.metadata.title=="پرونده ملک مادر P-0106"

    events=AuditEvent.objects.filter(action="MOTHER_PROPERTY_REPORT_EXPORT",entity_type="MotherProperty",entity_id="P-0106")
    assert events.count()==2
    assert {event.after["format"] for event in events}=={"XLSX","PDF"}
