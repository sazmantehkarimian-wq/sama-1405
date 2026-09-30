import hashlib,zipfile
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).parents[1]/'authority/inputs/1405-07-06'
def test_checksums_match_manifest():
 expected={line.split(maxsplit=1)[1].lstrip('*'):line.split(maxsplit=1)[0] for line in (ROOT/'SHA256SUMS.txt').read_text().splitlines() if line.strip()}
 for name,digest in expected.items(): assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==digest
def test_workbook_baseline_is_derived_from_rows(tmp_path):
 package=ROOT/'بسته_به_روزرسانی_سه_اکسل_سما_6مهر.zip'
 with zipfile.ZipFile(package) as z:
  for m in z.infolist():
   name=m.filename.encode('cp437').decode('utf8');(tmp_path/name).write_bytes(z.read(m))
 counts={}
 for p in tmp_path.glob('*.xlsx'):
  wb=load_workbook(p,read_only=True,data_only=True)
  if 'املاک مادر' in wb.sheetnames: key='mother';sheet='املاک مادر'
  else:key='active' if 'فعال' in wb['00_شروع']['A1'].value else 'inactive';sheet='فضاها'
  codes=[str(r[0]).strip() for r in wb[sheet].iter_rows(min_row=5,values_only=True) if r and r[0] not in (None,'')]
  counts[key]=(len(codes),len(set(codes)))
 assert counts=={'mother':(225,225),'active':(350,350),'inactive':(151,151)}
 assert counts['active'][1]+counts['inactive'][1]==501
