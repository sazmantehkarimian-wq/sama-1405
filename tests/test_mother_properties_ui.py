import io
import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from domains.properties.models import MotherProperty,CommercialSpace,MotherPropertySpaceLink,Region
from domains.registry.models import ImportBatch,SourceFile
pytestmark=pytest.mark.django_db

def records():
 region=Region.objects.create(code='R1',name='منطقه یک');prop=MotherProperty.objects.create(identifier='P-0001',name='ملک مرجع',region=region,primary_usage='فرهنگی',area=100,source_row=2);space=CommercialSpace.objects.create(code='S-1',name='فضای مرتبط',status='ACTIVE',region=region,current_usage='فرهنگی',area=20,source_row=3,source_classification='authority');batch=ImportBatch.objects.create(source_package='authority.zip',package_sha256='a'*64);source=SourceFile.objects.create(batch=batch,filename='mother.xlsx',sha256='b'*64,byte_size=1);MotherPropertySpaceLink.objects.create(mother_property=prop,space=space,evidence='ردیف صریح شیت فضاهای مرتبط',source_file=source,source_sheet='فضاهای مرتبط',source_row=2);return prop,space

def test_mother_list_dossier_cross_links_kpi_and_active_nav(client):
 prop,space=records();user=get_user_model().objects.create_user('u');client.force_login(user)
 dashboard=client.get('/').content.decode();assert f'href="/mother-properties/"' in dashboard and 'املاک مادر' in dashboard
 listing=client.get('/mother-properties/',{'identifier':'P-0001'});assert listing.status_code==200 and 'ملک مرجع' in listing.content.decode() and 'aria-current="page"' in listing.content.decode()
 dossier=client.get('/mother-properties/P-0001/').content.decode();assert 'فضای مرتبط' in dossier and f'/spaces/{space.code}/' in dossier and 'ردیف صریح' in dossier
 space_body=client.get(f'/spaces/{space.code}/').content.decode();assert '/mother-properties/P-0001/' in space_body

def test_mother_report_exact_layout_blank_columns_and_rtl(client):
 records();user=get_user_model().objects.create_user('u');client.force_login(user)
 params=[('layout','field:identifier'),('layout','blank:توضیحات کارشناس'),('layout','field:name'),('layout','blank:امضاء'),('field','identifier'),('field','name'),('blank','توضیحات کارشناس'),('blank','امضاء'),('sort','region'),('sort','identifier')]
 response=client.get('/mother-properties/report.xlsx',params);assert response.status_code==200
 ws=load_workbook(io.BytesIO(response.content)).active
 assert ws.sheet_view.rightToLeft is True and [c.value for c in ws[4]]==['شناسه ملک','توضیحات کارشناس','نام ملک / مرکز','امضاء'] and ws.freeze_panes=='A5' and ws.auto_filter.ref

def test_mother_filter_does_not_infer_identifier_equivalence(client):
 MotherProperty.objects.create(identifier='P-0001',name='الف',source_row=2);MotherProperty.objects.create(identifier='MP-000001',name='ب',source_row=3);user=get_user_model().objects.create_user('u');client.force_login(user)
 body=client.get('/mother-properties/',{'identifier':'P-0001'}).content.decode();assert 'الف' in body and '>ب<' not in body

def test_mother_pdf_and_docx_are_official_rtl_outputs(client):
 import zipfile
 from pypdf import PdfReader
 prop,_=records();user=get_user_model().objects.create_user('u2');client.force_login(user)
 query={'field':['identifier','name'],'blank':['امضاء']}
 pdf_response=client.get('/mother-properties/report.pdf',query);reader=PdfReader(io.BytesIO(pdf_response.content));text=''.join(page.extract_text() or '' for page in reader.pages)
 assert 'سما' not in text and 'SMK' not in text and len(reader.pages)>=1
 docx_response=client.get('/mother-properties/report.docx',query)
 with zipfile.ZipFile(io.BytesIO(docx_response.content)) as archive:
  document=archive.read('word/document.xml').decode();header=b''.join(archive.read(name) for name in archive.namelist() if name.startswith('word/header')).decode(errors='ignore')
 assert '<w:bidiVisual' in document and '<w:bidi' in document and 'SMK' not in document+header and 'SAMA' not in document+header
