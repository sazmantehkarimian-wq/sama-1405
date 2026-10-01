import io

import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from pypdf import PdfReader

from domains.contracts.models import Beneficiary, BeneficiaryAssignment, Contract
from domains.documents.models import Document
from domains.identity.models import AuditEvent
from domains.operations.models import Appraisal, Appraiser
from domains.properties.models import CommercialSpace


@pytest.fixture
def scoped_user(db):
    return get_user_model().objects.create_user("scoped-report", password="A-very-safe-password")


@pytest.mark.django_db
def test_scoped_search_distinguishes_not_found_from_empty_domain(client, scoped_user):
    client.force_login(scoped_user)
    CommercialSpace.objects.create(code="173", name="فضای ۱۷۳", status="ACTIVE")

    missing = client.get("/reports/space/", {"code": "999"})
    assert missing.status_code == 200
    assert "کد فضا پیدا نشد" in missing.content.decode()

    empty = client.get("/reports/space/", {"code": "173", "domain": "contracts"})
    assert empty.status_code == 200
    body = empty.content.decode()
    assert "فضای ۱۷۳" in body
    assert "کد فضا پیدا شد، اما در این حوزه داده‌ای ثبت نشده است" in body


@pytest.mark.django_db
def test_scoped_search_loads_only_requested_sections_and_counts_records(client, scoped_user):
    client.force_login(scoped_user)
    space = CommercialSpace.objects.create(code="174", name="فضای ۱۷۴", status="ACTIVE")
    beneficiary = Beneficiary.objects.create(kind="NATURAL", name="علی نمونه", first_name="علی", last_name="نمونه")
    BeneficiaryAssignment.objects.create(
        space=space, beneficiary=beneficiary, role="بهره‌بردار",
        start_date="1405/01/01", end_date="", status="ACTIVE",
    )
    Contract.objects.create(
        space=space, beneficiary=beneficiary, number="C-174",
        start_date="1405/01/01", end_date="1405/12/29", amount_rial=1000000,
    )

    response = client.get("/reports/space/", {"code": "174", "domain": ["contracts", "beneficiaries"]})
    assert response.status_code == 200
    sections = response.context["sections"]
    assert [section["key"] for section in sections] == ["contracts", "beneficiaries"]
    assert [section["count"] for section in sections] == [1, 1]
    body = response.content.decode()
    assert "C-174" in body and "علی نمونه" in body
    assert "کارشناسی" not in [section["title"] for section in sections]


@pytest.mark.django_db
def test_full_scoped_report_exports_multi_sheet_excel_and_pdf_with_official_identity(client, scoped_user, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    client.force_login(scoped_user)
    space = CommercialSpace.objects.create(code="175", name="فضای کامل", status="ACTIVE")
    expert = Appraiser.objects.create(first_name="مریم", last_name="ارزیاب")
    Appraisal.objects.create(
        space=space, appraiser_ref=expert, appraiser=expert.full_name,
        appraisal_date="1405/06/01", amount_rial=2500000, is_current=True,
    )

    xlsx = client.get("/reports/space.xlsx", {"code": "175", "domain": ["space", "appraisals"]})
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.content), data_only=True)
    assert len(workbook.sheetnames) == 2
    assert all(workbook[name].sheet_view.rightToLeft for name in workbook.sheetnames)
    all_values = [cell.value for name in workbook.sheetnames for row in workbook[name].iter_rows() for cell in row if cell.value is not None]
    assert "سازمان فرهنگی هنری شهرداری تهران" in all_values
    assert "175" in [str(value) for value in all_values]
    assert "سامانه مدیریت قراردادها" not in all_values

    pdf = client.get("/reports/space.pdf", {"code": "175", "domain": ["space", "appraisals"]})
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    reader = PdfReader(io.BytesIO(pdf.content))
    assert len(reader.pages) >= 1
    assert reader.metadata.title == "پرونده کد فضا 175"

    events = AuditEvent.objects.filter(action="SCOPED_REPORT_EXPORT", entity_type="CommercialSpace", entity_id="175")
    assert events.count() == 2
    assert {event.after["format"] for event in events} == {"XLSX", "PDF"}


@pytest.mark.django_db
def test_scoped_documents_include_archived_history(client, scoped_user, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    client.force_login(scoped_user)
    space = CommercialSpace.objects.create(code="176", name="فضای سند", status="ACTIVE")
    from django.core.files.uploadedfile import SimpleUploadedFile
    from services.documents import store_document, archive_document

    document = store_document(
        uploaded=SimpleUploadedFile("letter.pdf", b"%PDF-1.4\nSAMA", content_type="application/pdf"),
        title="نامه قدیمی", document_type="نامه", entity_type="CommercialSpace", entity_id=space.code,
        user=scoped_user, reference="R-176", document_date="1405/05/01",
    )
    archive_document(document=document, user=scoped_user, reason="جایگزینی نسخه")

    response = client.get("/reports/space/", {"code": "176", "domain": "documents"})
    assert response.status_code == 200
    body = response.content.decode()
    assert "نامه قدیمی" in body
    assert "بایگانی‌شده" in body
