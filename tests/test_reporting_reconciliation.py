import io, zipfile
import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from domains.properties.models import CommercialSpace, MotherProperty, Region

@pytest.fixture
def user(db):return get_user_model().objects.create_user('reconcile',password='safe-password')

def _data_rows(sheet,header_label):
 header=next(row for row in range(1,sheet.max_row+1) if sheet.cell(row,1).value==header_label)
 return sheet.max_row-header

@pytest.mark.django_db
def test_mother_export_ignores_pagination_and_respects_filter(client,user):
 region=Region.objects.create(code='4',name='منطقه ۴')
 for index in range(30):MotherProperty.objects.create(identifier=f'M-{index:03}',name=f'ملک {index}',region=region if index<12 else None,source_row=index+1)
 client.force_login(user)
 page=client.get('/mother-properties/');assert len(page.context['page'].object_list)==25 and page.context['total']==30
 sheet=load_workbook(io.BytesIO(client.get('/mother-properties/report.xlsx').content)).active
 assert _data_rows(sheet,'شناسه ملک')==30
 filtered=client.get('/mother-properties/',{'region_id':region.pk});filtered_sheet=load_workbook(io.BytesIO(client.get('/mother-properties/report.xlsx',{'region_id':region.pk}).content)).active
 assert filtered.context['total']==12 and _data_rows(filtered_sheet,'شناسه ملک')==12

@pytest.mark.django_db
def test_region_kpi_full_list_and_all_formats_reconcile(client,user):
 region=Region.objects.create(code='4',name='منطقه ۴')
 for index in range(23):CommercialSpace.objects.create(code=str(index+1),name=f'فضا {index+1}',region=region,status='ACTIVE' if index<17 else 'OUT_OF_CYCLE',source_row=index+1,source_classification='authority')
 client.force_login(user);page=client.get(f'/regions/{region.pk}/')
 assert page.context['space_count']==23 and page.context['active_count']==17 and len(page.context['spaces'])==23
 for kind,count in (('all',23),('active',17)):
  sheet=load_workbook(io.BytesIO(client.get(f'/regions/{region.pk}/reports/{kind}.xlsx').content)).active
  assert _data_rows(sheet,'کد فضا')==count
  assert client.get(f'/regions/{region.pk}/reports/{kind}.pdf').status_code==200
  docx=client.get(f'/regions/{region.pk}/reports/{kind}.docx').content
  with zipfile.ZipFile(io.BytesIO(docx)) as archive:
   xml=archive.read('word/document.xml').decode(errors='ignore');assert xml.count('<w:tr>')-1==count
 html=client.get(f'/regions/{region.pk}/reports/all.print').content.decode()
 assert 'مشاهده' not in html and 'SMK' not in html and 'سما' not in html and 'localhost' not in html

@pytest.mark.django_db
def test_dossier_critical_empty_states_are_explicit(client,user):
 space=CommercialSpace.objects.create(code='437',name='نمونه',status='ACTIVE',source_row=1,source_classification='authority')
 client.force_login(user);html=client.get('/spaces/437/').content.decode()
 for text in ('بهره‌بردار فعلی قطعی ثبت نشده','کارشناسی معتبر ثبت نشده','ملک مادر مرتبط قطعی ثبت نشده','گردش پرونده باز ثبت نشده','قرارداد جاری ثبت نشده'):assert text in html
