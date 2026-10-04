from pathlib import Path
from io import BytesIO
import pytest
from django.contrib.auth import get_user_model
from django.test import RequestFactory
from openpyxl import load_workbook
from domains.contracts.models import Contract
from domains.operations.models import Alert, Appraisal, WorkflowInstance
from domains.properties.models import Center, CommercialSpace, MotherProperty, Region
from queries.kpis import dashboard_kpis, scoped_counts
from reporting.engine import excel

@pytest.mark.django_db
def test_dashboard_kpi_counts_equal_their_drilldowns(client):
 user=get_user_model().objects.create_user('kpi-user',password='x');client.force_login(user)
 mother=MotherProperty.objects.create(identifier='MP-1',name='ملک',source_row=1)
 active=CommercialSpace.objects.create(code='1',name='فعال',status='ACTIVE',source_row=1,source_classification='authority')
 CommercialSpace.objects.create(code='2',name='خارج',status='OUT_OF_CYCLE',source_row=2,source_classification='authority')
 Appraisal.objects.create(space=active,appraisal_date='1405/01/01',appraiser='کارشناس',amount_rial=0,status='ثبت')
 dashboard=client.get('/');assert dashboard.status_code==200
 items={item.key:item for item in dashboard.context['kpis']}
 assert items['properties'].count==1 and client.get(items['properties'].url).context['total']==1
 assert items['active'].count==1 and client.get(items['active'].url).context['total']==1
 assert items['inactive'].count==1 and client.get(items['inactive'].url).context['total']==1
 assert items['without_appraisal'].count==0 and client.get(items['without_appraisal'].url).context['total']==0
 assert items['without_contract'].count==1 and client.get(items['without_contract'].url).context['total']==1

@pytest.mark.django_db
def test_region_scope_excludes_special_centers_and_center_scope_is_exact(client):
 user=get_user_model().objects.create_user('scope-user',password='x');client.force_login(user)
 region=Region.objects.create(code='14',name='منطقه ۱۴');regular=Center.objects.create(name='مرکز عادی',region=region);special=Center.objects.create(name='مرکز خاص',region=region,is_special=True)
 normal=CommercialSpace.objects.create(code='10',name='فضای عادی',status='ACTIVE',region=region,center=regular,source_row=1,source_classification='authority')
 exceptional=CommercialSpace.objects.create(code='11',name='فضای خاص',status='ACTIVE',region=region,center=special,source_row=2,source_classification='authority')
 region_page=client.get(f'/regions/{region.pk}/');center_page=client.get(f'/regions/special/{special.pk}/')
 assert region_page.context['space_count']==1 and list(region_page.context['spaces'])==[normal]
 assert center_page.context['space_count']==1 and list(center_page.context['spaces'])==[exceptional]
 assert 'region-rail' not in region_page.content.decode()

@pytest.mark.django_db
def test_blank_rows_are_output_only_and_saved_engine_is_rtl():
 before=CommercialSpace.objects.count()
 space=CommercialSpace.objects.create(code='5',name='نمونه',status='ACTIVE',source_row=1,source_classification='authority')
 payload=excel([space],selected=['code','name'],blank_rows=5)
 sheet=load_workbook(BytesIO(payload)).active
 assert sheet.sheet_view.rightToLeft and sheet.max_row==10
 assert all(sheet.cell(row,1).value is None for row in range(6,11))
 assert CommercialSpace.objects.count()==before+1

def test_rejected_visual_patterns_are_absent():
 templates='\n'.join(path.read_text() for path in Path('ui/templates/ui').glob('*.html'))
 assert 'region-rail' not in templates and 'sidebar' not in templates
 css=Path('design_system/static/design_system/css/app.css').read_text()
 assert '.region-grid' in css and '.center-grid' in css and '@media print' in css
