import io,zipfile,pytest,hashlib
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from domains.properties.models import CommercialSpace
from reporting.engine import excel,docx,pdf
from pypdf import PdfReader
from domains.identity.models import SavedReport, ArchivedReportSnapshot, AuditEvent
@pytest.fixture
def user(db):return get_user_model().objects.create_user('u',password='A-very-safe-password')
@pytest.mark.django_db
def test_space_list_and_dossier(client,user):
 client.force_login(user);s=CommercialSpace.objects.create(code='9',name='فضای آزمون',status='ACTIVE',source_row=5,source_classification='authority')
 assert client.get('/spaces/').status_code==200
 body=client.get('/spaces/9/').content.decode();assert 'پرونده فضای 9' in body and 'الان دست کیه؟' in body and 'None' not in body

@pytest.mark.django_db
def test_advanced_space_filters_and_saved_view(client,user):
 from domains.properties.models import Region
 region=Region.objects.create(name='منطقه ۱')
 CommercialSpace.objects.create(code='A-10',name='فضای بزرگ',status='ACTIVE',area=100,region=region,address='تهران',source_row=1,source_classification='authority')
 CommercialSpace.objects.create(code='B-20',name='فضای کوچک',status='OUT_OF_CYCLE',area=20,source_row=2,source_classification='authority')
 client.force_login(user)
 response=client.get('/spaces/',{'code':'A','code_op':'starts','status':['ACTIVE'],'area_min':'50','address_presence':'nonempty'})
 assert response.status_code==200 and 'A-10' in response.content.decode() and 'B-20' not in response.content.decode()
 response=client.post('/spaces/filters/save/',{'name':'فضاهای منتخب','status':['ACTIVE'],'area_min':'50'})
 assert response.status_code==302
 from domains.identity.models import SavedFilter
 saved=SavedFilter.objects.get(owner=user);assert saved.definition=={'status':['ACTIVE'],'area_min':['50']}
 assert 'status=ACTIVE' in client.get(f'/spaces/filters/{saved.pk}/').url
 response=client.get('/spaces/',{'code':'A-10','code_op':'equals','status':['OUT_OF_CYCLE'],'logic':'or'})
 body=response.content.decode();assert 'A-10' in body and 'B-20' in body
 response=client.get('/spaces/',{'status':'ACTIVE','contract_presence':'empty','appraisal_presence':'empty'})
 assert 'A-10' in response.content.decode() and 'B-20' not in response.content.decode()
 saved.name='نام قدیم';saved.save()
 assert client.post(f'/spaces/filters/{saved.pk}/rename/',{'name':'نام جدید'}).status_code==302
 saved.refresh_from_db();assert saved.name=='نام جدید'
 assert client.post(f'/spaces/filters/{saved.pk}/delete/').status_code==302
 assert not SavedFilter.objects.filter(pk=saved.pk).exists()
@pytest.mark.django_db
def test_real_export_structures():
 s=CommercialSpace.objects.create(code='9',name='فضا',status='ACTIVE',source_row=5,source_classification='authority');qs=CommercialSpace.objects.all()
 x=load_workbook(io.BytesIO(excel(qs)));ws=x.active;assert ws.sheet_view.rightToLeft and ws.auto_filter.ref
 assert [ws.cell(4,column).value for column in range(1,7)]==['کد فضا','نام فضا / مرکز','وضعیت','منطقه','کاربری','مساحت (مترمربع)']
 assert ws.freeze_panes=='A5' and ws.print_title_rows=='$4:$4' and ws.print_area
 assert ws['A4'].alignment.horizontal=='center' and ws['A4'].alignment.readingOrder==2
 with zipfile.ZipFile(io.BytesIO(docx(qs))) as z:
  assert 'word/document.xml' in z.namelist()
  assert any(name.startswith('word/media/') for name in z.namelist())
  xml=z.read('word/document.xml').decode();assert 'w:bidiVisual' in xml and 'w:bidi' in xml
  combined=''.join(z.read(name).decode(errors='ignore') for name in z.namelist() if name.endswith('.xml'))
  assert 'سازمان فرهنگی هنری شهرداری تهران' in combined and 'مدیریت اقتصادی و املاک' in combined
  assert 'سامانه مدیریت قراردادها' not in combined and 'SMK' not in combined
 pdf_payload=pdf(qs);assert pdf_payload.startswith(b'%PDF')
 rendered=PdfReader(io.BytesIO(pdf_payload));assert len(rendered.pages)>=1
 assert rendered.metadata.title=='گزارش رسمی املاک'
 extracted=''.join(page.extract_text() or '' for page in rendered.pages)
 assert 'http://' not in extracted and 'SMK' not in extracted and 'سامانه مدیریت قراردادها' not in extracted

@pytest.mark.django_db
def test_typed_operational_list_has_official_excel_export(client,user):
 client.force_login(user)
 response=client.get('/records/utilities/export.xlsx')
 assert response.status_code==200
 workbook=load_workbook(io.BytesIO(response.content))
 assert workbook.active.sheet_view.rightToLeft
 assert workbook.active.auto_filter.ref
 assert any(image for image in workbook.active._images)
 assert workbook.active.page_setup.orientation=='landscape'
 assert workbook.active.print_title_rows

@pytest.mark.django_db
def test_saved_report_and_immutable_snapshot(client,user,settings,tmp_path):
 settings.MEDIA_ROOT=tmp_path
 CommercialSpace.objects.create(code='501',name='فضای مرجع',status='ACTIVE',source_row=2,source_classification='authority')
 client.force_login(user)
 response=client.post('/reports/save/',{'name':'فضاهای فعال','status':'ACTIVE','field':['code','status'],'blank':['اقدام']})
 assert response.status_code==302
 report=SavedReport.objects.get(owner=user)
 assert report.filters=={'status':['ACTIVE']} and report.fields==['code','status']
 assert client.get(f'/reports/{report.pk}/open/').url=='/spaces/?status=ACTIVE&field=code&field=status&blank=%D8%A7%D9%82%D8%AF%D8%A7%D9%85'
 assert client.post(f'/reports/{report.pk}/archive/').status_code==302
 snapshot=ArchivedReportSnapshot.objects.get(report=report);payload=snapshot.file.read()
 assert snapshot.row_count==1 and hashlib.sha256(payload).hexdigest()==snapshot.sha256
 assert snapshot.query_context['filters']=={'status':['ACTIVE']}
 assert AuditEvent.objects.filter(action='REPORT_DEFINITION_CREATE').exists()
 assert AuditEvent.objects.filter(action='REPORT_SNAPSHOT_CREATE').exists()

@pytest.mark.django_db
def test_dossier_is_read_only_and_specialist_operations_preserve_space_context(client,user):
 s=CommercialSpace.objects.create(code='CTX-1',name='فضای زمینه',status='ACTIVE',source_row=1,source_classification='authority')
 client.force_login(user);response=client.get(f'/spaces/{s.code}/');body=response.content.decode()
 for forbidden in ('action="/spaces/CTX-1/contracts/"','action="/spaces/CTX-1/appraisals/"','action="/spaces/CTX-1/utilities/"','action="/spaces/CTX-1/documents/"','action="/spaces/CTX-1/movement/"','action="/spaces/CTX-1/alerts/"'):
  assert forbidden not in body
 for action,label in (('contract','ثبت قرارداد جدید'),('appraisal','ثبت کارشناسی جدید'),('document','بارگذاری سند'),('utility','ثبت انشعاب / مصرف'),('movement','ثبت تحویل'),('alert','ثبت مورد پیگیری')):
  assert label in body
  operation=client.get(f'/operations/{action}/?space={s.code}')
  assert operation.status_code==200 and f'پرونده فضای {s.code}' in operation.content.decode()
  assert f'/spaces/{s.code}/' in operation.content.decode()

@pytest.mark.django_db
def test_specialist_module_entry_without_context_requires_structured_space_selection(client,user):
 CommercialSpace.objects.create(code='SELECT-1',name='فضای انتخاب',status='ACTIVE',source_row=1,source_classification='authority')
 client.force_login(user)
 for domain,action in (('contracts','contract'),('appraisals','appraisal'),('utilities','utility'),('documents','document'),('alerts','alert')):
  listing=client.get(f'/records/{domain}/').content.decode();assert f'/operations/{action}/' in listing
  page=client.get(f'/operations/{action}/').content.decode();assert 'انتخاب فضای تجاری' in page and 'SELECT-1' in page
