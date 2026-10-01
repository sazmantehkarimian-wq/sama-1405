from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook

from domains.operations.models import UtilityBill, UtilityConnection, UtilityMeasurement
from domains.properties.models import CommercialSpace, Region


@pytest.mark.django_db
def test_utility_dashboard_filters_water_gas_and_measurement(client):
    user=get_user_model().objects.create_user("utility-report-user",password="A-very-safe-password")
    client.force_login(user)
    region=Region.objects.create(name="منطقه آزمون")
    water_space=CommercialSpace.objects.create(code="8101",name="فضای آب",status="ACTIVE",region=region)
    gas_space=CommercialSpace.objects.create(code="8102",name="فضای گاز",status="ACTIVE",region=region)
    water=UtilityConnection.objects.create(space=water_space,utility_type="WATER",account_number="W-100",status="ACTIVE")
    gas=UtilityConnection.objects.create(space=gas_space,utility_type="GAS",account_number="G-200",status="ACTIVE")
    measurement=UtilityMeasurement.objects.create(
        space=water_space,utility_type="WATER",period_start="1405/01/01",period_end="1405/01/31",
        consumption="12.500",reading_date="1405/01/31",measurement_unit="m3",is_valid=True,
    )
    UtilityBill.objects.create(
        connection=water,period_start="1405/01/01",period_end="1405/01/31",
        amount_rial=1000,consumption="12.500",measurement=measurement,payment_status="PAID",payment_date="1405/02/01",
    )
    UtilityBill.objects.create(
        connection=gas,period_start="1405/01/01",period_end="1405/01/31",
        amount_rial=2000,consumption="20.000",payment_status="UNPAID",
    )

    response=client.get("/utilities/",{"utility_type":"WATER","measurement":"yes"})
    assert response.status_code==200
    body=response.content.decode()
    assert "W-100" in body and "G-200" not in body
    assert response.context["bill_count"]==1
    assert response.context["total_amount"]==1000

    response=client.get("/utilities/",{"q":"8102"})
    body=response.content.decode()
    assert "G-200" in body and "W-100" not in body


@pytest.mark.django_db
def test_utility_dashboard_period_comparison_is_per_connection(client):
    user=get_user_model().objects.create_user("utility-compare",password="A-very-safe-password")
    client.force_login(user)
    space=CommercialSpace.objects.create(code="8103",name="فضا",status="ACTIVE")
    connection=UtilityConnection.objects.create(space=space,utility_type="WATER",account_number="W-COMP",status="ACTIVE")
    UtilityBill.objects.create(connection=connection,period_start="1405/01/01",period_end="1405/01/31",amount_rial=1000,consumption="10",payment_status="UNKNOWN")
    UtilityBill.objects.create(connection=connection,period_start="1405/02/01",period_end="1405/02/31",amount_rial=1300,consumption="13",payment_status="UNKNOWN")

    response=client.get("/utilities/")
    assert response.status_code==200
    rows=list(response.context["page"].object_list)
    latest=next(item for item in rows if item.period_start=="1405/02/01")
    assert latest.previous_amount_rial==1000
    assert latest.amount_change_rial==300
    assert latest.previous_consumption==10
    assert latest.consumption_change==3


@pytest.mark.django_db
def test_utility_excel_export_respects_filters_and_is_excel_2019_compatible(client):
    user=get_user_model().objects.create_user("utility-excel",password="A-very-safe-password")
    client.force_login(user)
    water_space=CommercialSpace.objects.create(code="8104",name="فضای آب",status="ACTIVE")
    gas_space=CommercialSpace.objects.create(code="8105",name="فضای گاز",status="ACTIVE")
    water=UtilityConnection.objects.create(space=water_space,utility_type="WATER",account_number="WX",status="ACTIVE")
    gas=UtilityConnection.objects.create(space=gas_space,utility_type="GAS",account_number="GX",status="ACTIVE")
    UtilityBill.objects.create(connection=water,period_start="1405/01/01",period_end="1405/01/31",amount_rial=1100,payment_status="UNPAID")
    UtilityBill.objects.create(connection=gas,period_start="1405/01/01",period_end="1405/01/31",amount_rial=2200,payment_status="UNPAID")

    response=client.get("/utilities/export.xlsx",{"utility_type":"GAS"})
    assert response.status_code==200
    assert response["Content-Type"]=="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook=load_workbook(BytesIO(response.content),data_only=True)
    sheet=workbook.active
    values=[cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert "GX" in values
    assert "WX" not in values
    assert sheet.sheet_view.rightToLeft is True
