import zipfile
from pathlib import Path
import pytest
from django.contrib.auth import get_user_model
from django.test import Client
import io
from openpyxl import load_workbook
from import_pipeline.import_authorities import first_value, run
from domains.properties.models import MotherProperty, CommercialSpace
from domains.registry.models import RawCell, CanonicalField, Discrepancy
from domains.operations.models import Appraisal, Auction, DecisionOrder, UtilityObligation, SourceDocumentReference, TimelineEvent
from domains.contracts.models import Beneficiary, Contract

ROOT=Path(__file__).parents[1]/'authority/inputs/1405-07-06'

def test_alternate_contract_and_appraisal_headings_are_not_dropped():
 row={'مبلغ با ارزش افزوده (ریال)':120,'مبلغ کارشناسی (ریال)':80,'نام کارشناس':'کارشناس دوم','مرجع نامه':'نامه ۲'}
 assert first_value(row,'مبلغ قرارداد با ارزش افزوده (ریال)','مبلغ با ارزش افزوده (ریال)')==120
 assert first_value(row,'مبلغ کارشناسی اجاره ماهانه (ریال)','مبلغ کارشناسی (ریال)')==80
 assert first_value(row,'نام کارشناس ارزیابی','نام کارشناس')=='کارشناس دوم'
 assert first_value(row,'شماره / مرجع نامه کارشناسی','مرجع نامه')=='نامه ۲'

@pytest.mark.django_db(transaction=True)
def test_complete_authority_import_is_lossless_and_typed(tmp_path):
 package=tmp_path/'بسته_به_روزرسانی_سه_اکسل_سما_6مهر.zip'
 package.write_bytes((ROOT/package.name).read_bytes())
 with zipfile.ZipFile(package) as archive:
  expected_nonempty=0
  for member in archive.infolist():
   name=member.filename.encode('cp437').decode('utf8')
   target=tmp_path/name; target.write_bytes(archive.read(member))
   workbook=load_workbook(target,read_only=True,data_only=True)
   expected_nonempty += sum(1 for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row if value not in (None,''))
 result=run(package)
 assert result['mother_properties']==225
 assert result['active_spaces']==350
 assert result['out_of_cycle_spaces']==151
 assert result['unique_spaces']==501
 assert MotherProperty.objects.values('identifier').distinct().count()==225
 assert CommercialSpace.objects.values('code').distinct().count()==501
 assert RawCell.objects.count()==expected_nonempty
 assert not CanonicalField.objects.filter(persian_label__regex=r'^ستون [0-9]+$').exists()
 assert Discrepancy.objects.filter(entity_type='CommercialSpace').exists()
 for model in (Appraisal,Auction,DecisionOrder,UtilityObligation,SourceDocumentReference,TimelineEvent):
  assert model.objects.exists(), model.__name__
 # Representative real-row semantic trace: displaced status text is retained as evidence, never as identity.
 active_one=CommercialSpace.objects.get(code='1')
 contract=active_one.contracts.get(start_date='1404/07/01')
 assert contract.number=='' and contract.status=='منعقدشده'
 assert Discrepancy.objects.filter(entity_key='1',field_key='contract.number',observed_value='در حال انعقاد قرارداد').exists()
 assert not CommercialSpace.objects.get(code='5').contracts.exists()
 assert Discrepancy.objects.filter(entity_key='5',field_key='contract.number',observed_value='خارج از مزایده').exists()
 assert not Beneficiary.objects.filter(name__in=['آماده مزایده','خارج از مزایده','فاقد بهره بردار','فروشگاه محصولات فرهنگی','جمع آوری شده','محصولات']).exists()
 assert Discrepancy.objects.filter(entity_key='320',field_key='beneficiary.name',observed_value='محصولات').exists()
 assert Discrepancy.objects.filter(entity_key='476',field_key='beneficiary.name',observed_value='جمع آوری شده').exists()
 assert Beneficiary.objects.filter(name='تندیس خسروی راد',kind='UNSPECIFIED').exists()
 appraisal=active_one.appraisals.get(year='1404')
 assert appraisal.appraisal_date=='1404/02/01' and appraisal.status==''
 assert appraisal.appraiser=='نوید دولت‌آبادی' and appraisal.amount_rial==85000000
 # Authority reconciliation covers valid/missing numbers and dates, formation status and expired history.
 assert Contract.objects.filter(space__code='3',number='145072',start_date='1404/07/01',end_date='1405/06/31',status='در حال انعقاد').exists()
 assert Contract.objects.filter(space__code='11',number='',start_date='',status='در حال انعقاد').exists()
 assert Contract.objects.filter(end_date__lt='1405/07/01').exclude(end_date='').exists()
 assert Contract.objects.count()==332
 assert CommercialSpace.objects.filter(status='ACTIVE',contracts__isnull=True).count()==70
 # Appraisal zero remains a known numeric zero; it is not normalized to an unknown/null value.
 assert Appraisal.objects.filter(space__code='92',year='1402',amount_rial=0,appraisal_date='',appraiser='',status='قابل استفاده با نقص اطلاعات').exists()
 assert Appraisal.objects.filter(amount_rial__isnull=True).exists()
 assert TimelineEvent.objects.filter(space=active_one,event_type='CONTRACT_HISTORY',source_entity_id=str(contract.pk)).exists()
 assert TimelineEvent.objects.filter(space=active_one,event_type='BENEFICIARY_HISTORY').exists()
 assert TimelineEvent.objects.filter(space=active_one,event_type='APPRAISAL_HISTORY',jalali_date='1404/02/01').exists()
 # Continue the trace through list, dossier and the official typed export.
 user=get_user_model().objects.create_user('trace-user',password='A-very-safe-password');client=Client();client.force_login(user)
 contract_response=client.get('/records/contracts/')
 assert contract_response.context['page'].paginator.count==332
 contract_list=contract_response.content.decode()
 assert client.get('/spaces/?status=ACTIVE&contract_presence=empty').context['total']==70
 dashboard=client.get('/');assert dashboard.context['contract_count']==332 and dashboard.context['without_contract_count']==70
 assert 'در حال انعقاد قرارداد' not in contract_list and 'خارج از مزایده' not in contract_list
 appraisal_list=client.get('/records/appraisals/?q=1').content.decode()
 assert '۱۴۰۴/۰۲/۰۱' in appraisal_list and '۸۵,۰۰۰,۰۰۰ ریال' in appraisal_list
 dossier=client.get('/spaces/1/').content.decode()
 assert 'نوید دولت‌آبادی' in dossier and '۱۴۰۴/۰۲/۰۱' in dossier
 exported=load_workbook(io.BytesIO(client.get('/records/contracts/export.xlsx').content)).active
 assert all(cell.value not in ('در حال انعقاد قرارداد','خارج از مزایده','آماده مزایده') for row in exported.iter_rows() for cell in row)
