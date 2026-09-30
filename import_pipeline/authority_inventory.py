"""Independent inventory and integrity checks for the 6 Mehr authority package."""
from __future__ import annotations

import hashlib
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path, PurePosixPath
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

EXPECTED_FILES = {
    "املاک_مادر_به‌روزشده_6مهر.xlsx": "mother",
    "فضاهای_تجاری_فعال_به‌روزشده_6مهر.xlsx": "active",
    "فضاهای_از_دور_خارج_شده_به‌روزشده_6مهر.xlsx": "out_of_cycle",
}
APPROVED_BASELINE = {"mother_properties": 225, "active_spaces": 350, "out_of_cycle_spaces": 151, "unique_spaces": 501}


@dataclass(frozen=True)
class AuthorityInventory:
    mother_properties: int
    active_spaces: int
    out_of_cycle_spaces: int
    unique_spaces: int
    overlap: tuple[str, ...]
    duplicate_mother_identifiers: tuple[str, ...]
    duplicate_active_codes: tuple[str, ...]
    duplicate_out_of_cycle_codes: tuple[str, ...]

    def as_dict(self):
        return asdict(self)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def verify_manifest(directory: Path) -> dict[str, str]:
    """Verify every manifest entry and reject traversal outside the authority directory."""
    verified = {}
    for line in (directory / "SHA256SUMS.txt").read_text(encoding="utf-8-sig").splitlines():
        if not line.strip():
            continue
        expected, filename = line.split(maxsplit=1)
        filename = filename.lstrip(" *")
        relative = PurePosixPath(filename)
        if relative.is_absolute() or ".." in relative.parts:
            raise ValueError(f"Unsafe manifest path: {filename}")
        actual = sha256(directory / filename)
        if actual != expected.lower():
            raise ValueError(f"SHA-256 mismatch for {filename}: {actual}")
        verified[filename] = actual
    return verified


def _normal(value) -> str:
    value = "" if value is None else str(value).strip()
    value = value.translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789"))
    return value[:-2] if value.endswith(".0") else value


def _identifiers(path: Path, sheet: str, heading: str) -> list[str]:
    worksheet = load_workbook(path, read_only=True, data_only=True)[sheet]
    heading_index = None
    values = []
    for row in worksheet.iter_rows(values_only=True):
        if heading_index is None:
            headings = [_normal(value) for value in row]
            if heading in headings:
                heading_index = headings.index(heading)
            continue
        if heading_index < len(row):
            value = _normal(row[heading_index])
            if value:
                values.append(value)
    if heading_index is None:
        raise ValueError(f"Heading {heading!r} not found in {path.name}/{sheet}")
    return values


def _duplicates(values: list[str]) -> tuple[str, ...]:
    seen, duplicates = set(), set()
    for value in values:
        (duplicates if value in seen else seen).add(value)
    return tuple(sorted(duplicates))


def inspect_package(package: Path) -> AuthorityInventory:
    """Read actual workbook rows; no database or display count is trusted."""
    with TemporaryDirectory(prefix="sama-authority-") as temporary:
        root = Path(temporary)
        extracted = {}
        with zipfile.ZipFile(package) as archive:
            for member in archive.infolist():
                try:
                    name = member.filename.encode("cp437").decode("utf-8")
                except (UnicodeEncodeError, UnicodeDecodeError):
                    name = member.filename
                basename = Path(name).name
                if basename in EXPECTED_FILES:
                    target = root / basename
                    target.write_bytes(archive.read(member))
                    extracted[basename] = target
        if set(extracted) != set(EXPECTED_FILES):
            missing = sorted(set(EXPECTED_FILES) - set(extracted))
            raise ValueError(f"Authority package members mismatch; missing: {missing}")

        by_kind = {kind: extracted[name] for name, kind in EXPECTED_FILES.items()}
        mother = _identifiers(by_kind["mother"], "املاک مادر", "شناسه ملک مادر")
        active = _identifiers(by_kind["active"], "فضاها", "کد فضای تجاری")
        inactive = _identifiers(by_kind["out_of_cycle"], "فضاها", "کد فضای تجاری")
        inventory = AuthorityInventory(
            mother_properties=len(mother), active_spaces=len(active), out_of_cycle_spaces=len(inactive),
            unique_spaces=len(set(active) | set(inactive)), overlap=tuple(sorted(set(active) & set(inactive))),
            duplicate_mother_identifiers=_duplicates(mother), duplicate_active_codes=_duplicates(active),
            duplicate_out_of_cycle_codes=_duplicates(inactive),
        )
        actual = {key: getattr(inventory, key) for key in APPROVED_BASELINE}
        defects = inventory.overlap + inventory.duplicate_mother_identifiers + inventory.duplicate_active_codes + inventory.duplicate_out_of_cycle_codes
        if actual != APPROVED_BASELINE or defects:
            raise ValueError(f"Authority inventory failed: counts={actual}, identity_defects={defects}")
        return inventory
