from __future__ import annotations
import hashlib, json, zipfile
from pathlib import Path
from decimal import Decimal,InvalidOperation
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from domains.registry.models import ImportBatch,SourceFile,CanonicalField,RawCell,Discrepancy
from domains.properties.models import Region,Center,MotherProperty,CommercialSpace,CommercialSpaceStatusHistory,MotherPropertySpaceLink
from domains.contracts.models import Beneficiary,BeneficiaryAssignment,Contract
from domains.operations.models import Appraisal,Auction,TimelineEvent,DecisionOrder,UtilityObligation,SourceDocumentReference
from services.dates import normalize_jalali
from import_pipeline.authority_inventory import inspect_package
MAPPED={
 'کد فضای تجاری':('space.code','کد فضای تجاری','string','properties','CommercialSpace'), 'نام فضا / مرکز':('space.name','نام فضا / مرکز','string','properties','CommercialSpace'),
 'منطقه شهرداری':('space.region','منطقه شهرداری','string','properties','CommercialSpace'), 'حوزه سازمانی':('space.scope','حوزه سازمانی','string','properties','CommercialSpace'),
 'وضعیت فضا':('space.status','وضعیت فضا','enum','properties','CommercialSpace'), 'نوع فضا / دارایی':('space.asset_type','نوع فضا / دارایی','string','properties','CommercialSpace'),
 'مساحت اعیان (مترمربع)':('space.area','مساحت اعیان','decimal','properties','CommercialSpace'), 'نشانی فضای تجاری':('space.address','نشانی فضای تجاری','text','properties','CommercialSpace'),
 'شرح موقعیت مکانی':('space.physical_details','شرح موقعیت مکانی','text','properties','CommercialSpace'), 'مشخصات فیزیکی و امکانات':('space.physical_details','مشخصات فیزیکی و امکانات','text','properties','CommercialSpace'),
 'کاربری اصلی فضای تجاری':('space.current_usage','کاربری اصلی','string','properties','CommercialSpace'), 'فعالیت / کاربری پیشنهادی':('space.proposed_activity','فعالیت پیشنهادی','string','properties','CommercialSpace'),
 'گروه فعالیت پیشنهادی':('space.activity_group','گروه فعالیت پیشنهادی','string','properties','CommercialSpace'), 'کاربری پیشین':('space.previous_usage','کاربری پیشین','string','properties','CommercialSpace'),
 'شناسه ملک مادر':('property.identifier','شناسه ملک مادر','string','properties','MotherProperty'), 'نام ملک / مجموعه':('property.name','نام ملک / مجموعه','string','properties','MotherProperty'),
 'شماره قرارداد':('contract.number','شماره قرارداد','string','contracts','Contract'), 'نام بهره‌بردار / شخص':('beneficiary.name','نام بهره‌بردار','string','contracts','Beneficiary'),
 'نوع نقش':('beneficiary.role','نقش بهره‌بردار','string','contracts','BeneficiaryAssignment'), 'وضعیت رابطه':('beneficiary.status','وضعیت رابطه بهره‌بردار','string','contracts','BeneficiaryAssignment'),
 'تاریخ شروع':('contract.start_date','تاریخ شروع قرارداد/رابطه','jalali_date','contracts','Contract'), 'تاریخ پایان':('contract.end_date','تاریخ پایان قرارداد/رابطه','jalali_date','contracts','Contract'),
 'تاریخ انعقاد':('contract.signed_date','تاریخ انعقاد قرارداد','jalali_date','contracts','Contract'), 'وضعیت قرارداد':('contract.status','وضعیت قرارداد','string','contracts','Contract'),
 'مبلغ قرارداد با ارزش افزوده (ریال)':('contract.amount_rial','مبلغ قرارداد با ارزش افزوده','money','contracts','Contract'), 'مبلغ با ارزش افزوده (ریال)':('contract.amount_rial','مبلغ قرارداد با ارزش افزوده','money','contracts','Contract'),
 'تعهد سرمایه‌گذاری (ریال)':('contract.investment_commitment_rial','تعهد سرمایه‌گذاری','money','contracts','Contract'), 'وضعیت نسخه قرارداد':('contract.signed_state','وضعیت نسخه قرارداد','string','contracts','Contract'),
 'سال':('appraisal.year','سال کارشناسی','year','operations','Appraisal'), 'نوبت':('appraisal.sequence','نوبت کارشناسی/مزایده','string','operations','Appraisal'),
 'مبلغ کارشناسی اجاره ماهانه (ریال)':('appraisal.amount_rial','مبلغ کارشناسی','money','operations','Appraisal'), 'مبلغ کارشناسی (ریال)':('appraisal.amount_rial','مبلغ کارشناسی','money','operations','Appraisal'),
 'نام کارشناس ارزیابی':('appraisal.appraiser','نام کارشناس','string','operations','Appraisal'), 'نام کارشناس':('appraisal.appraiser','نام کارشناس','string','operations','Appraisal'),
 'شماره / مرجع نامه کارشناسی':('appraisal.reference','مرجع کارشناسی','string','operations','Appraisal'), 'مرجع نامه':('appraisal.reference','مرجع کارشناسی','string','operations','Appraisal'),
 'تاریخ کارشناسی / گزارش':('appraisal.appraisal_date','تاریخ کارشناسی','jalali_date','operations','Appraisal'), 'وضعیت بررسی':('source.review_status','وضعیت بررسی منبع','string','provenance','RawCell'),
 'نوع رویداد':('timeline.event_type','نوع رویداد','string','operations','TimelineEvent'), 'تاریخ رویداد':('timeline.date','تاریخ رویداد','jalali_date','operations','TimelineEvent'),
 'تاریخ':('timeline.date','تاریخ رویداد','jalali_date','operations','TimelineEvent'), 'دسته‌بندی':('timeline.category','دسته‌بندی رویداد','string','operations','TimelineEvent'), 'شرح':('timeline.description','شرح رویداد','text','operations','TimelineEvent'),
 'نوع سند':('document.type','نوع سند','string','documents','SourceDocumentReference'), 'شماره / مرجع سند':('document.reference','مرجع سند','string','documents','SourceDocumentReference'),
 'شماره / مرجع':('document.reference','مرجع سند','string','documents','SourceDocumentReference'), 'وضعیت سند':('document.status','وضعیت سند','string','documents','SourceDocumentReference'), 'مسیر / لینک فایل':('document.source_path','مسیر منبع سند','string','documents','SourceDocumentReference'),
}
MAPPED.update({
 'وضعیت اقاله':('contract.termination_status','وضعیت اقاله','string','contracts','Contract'),
 'وضعیت تخلیه':('contract.vacating_status','وضعیت تخلیه','string','contracts','Contract'),
 'توضیح دوره مبلغ':('contract.amount_period_note','توضیح دوره مبلغ','text','contracts','Contract'),
 'توضیحات':('source.notes','توضیحات منبع','text','provenance','RawCell'),
 'نتیجه / وضعیت مزایده':('auction.result','نتیجه مزایده','string','operations','Auction'),
 'نتیجه / وضعیت':('auction.result','نتیجه مزایده','string','operations','Auction'),
 'مرحله فرایند مزایده':('auction.stage','مرحله مزایده','string','operations','Auction'),
 'مرحله فرایند':('auction.stage','مرحله مزایده','string','operations','Auction'),
 'نوع تصمیم':('decision.type','نوع تصمیم','string','operations','DecisionOrder'),
 'تاریخ تصمیم':('decision.date','تاریخ تصمیم','jalali_date','operations','DecisionOrder'),
 'شماره نامه':('decision.letter_number','شماره نامه','string','operations','DecisionOrder'),
 'مرجع جلسه':('decision.session_reference','مرجع جلسه','string','operations','DecisionOrder'),
 'متن تصمیم / مصوبه':('decision.text','متن تصمیم یا مصوبه','text','operations','DecisionOrder'),
 'نوع انشعاب':('utility.type','نوع انشعاب','string','operations','UtilityObligation'),
 'نوع محاسبه':('utility.calculation_type','نوع محاسبه','string','operations','UtilityObligation'),
 'مبلغ (ریال)':('utility.amount_rial','مبلغ','money','operations','UtilityObligation'),
 'درصد':('utility.percentage','درصد','decimal','operations','UtilityObligation'),
 'توضیح استاندارد':('utility.standard_description','توضیح استاندارد','text','operations','UtilityObligation'),
 'نوع ایراد':('discrepancy.field','نوع مغایرت','string','registry','Discrepancy'),
 'مقدار فعلی':('discrepancy.observed_value','مقدار موجود','string','registry','Discrepancy'),
 'مقدار استاندارد / مقصد':('discrepancy.expected_value','مقدار مورد انتظار','string','registry','Discrepancy'),
 'شرح کنترل':('discrepancy.reason','شرح کنترل','text','registry','Discrepancy'),
 'وضعیت':('discrepancy.status','وضعیت مغایرت','string','registry','Discrepancy'),
})
INVALID_CONTRACT_NUMBERS={'در حال انعقاد قرارداد','خارج از مزایده','آماده مزایده','سامانه','قرارداد دارد','در دست اقدام','در دست انعقاد','تخلیه','فاقد قرارداد'}
INVALID_BENEFICIARIES={'آماده مزایده','خارج از مزایده','فاقد بهره بردار','فاقد بهره‌بردار','در حال انعقاد قرارداد','در دست اقدام','تخلیه','فروشگاه محصولات فرهنگی'}
def sha(path):
 h=hashlib.sha256()
 with open(path,'rb') as f:
  for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
 return h.hexdigest()
def text(v): return '' if v is None else str(v).strip()
def first_value(row,*headings):
 for heading in headings:
  if row.get(heading) not in (None,''):return row.get(heading)
 return None
def normalized_date(value):
 try:return normalize_jalali(value)
 except (ValueError,TypeError):return ''
def decimal(v):
 if v in (None,''): return None
 try: return Decimal(str(v).replace(',','').translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')))
 except InvalidOperation:return None
def norm_code(v):
 s=text(v).translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789'))
 return s[:-2] if s.endswith('.0') else s
def heading_row(ws):
 for n,row in enumerate(ws.iter_rows(min_row=1,max_row=12,values_only=True),1):
  vals=[text(x) for x in row]
  if 'کد فضای تجاری' in vals or 'شناسه ملک مادر' in vals: return n,vals
 return None,[]
OPERATIONAL_HEADINGS={
 'نام بهره‌بردار / شخص','نوع نقش','تاریخ شروع','تاریخ پایان','وضعیت رابطه','توضیحات','وضعیت بررسی',
 'تاریخ انعقاد','وضعیت قرارداد','مبلغ قرارداد با ارزش افزوده (ریال)','مبلغ با ارزش افزوده (ریال)',
 'وضعیت اقاله','وضعیت تخلیه','تعهد سرمایه‌گذاری (ریال)','وضعیت نسخه قرارداد','توضیح دوره مبلغ',
 'سال','نوبت','مبلغ کارشناسی اجاره ماهانه (ریال)','مبلغ کارشناسی (ریال)','نام کارشناس ارزیابی',
 'نام کارشناس','شماره / مرجع نامه کارشناسی','مرجع نامه','تاریخ کارشناسی / گزارش',
 'نتیجه / وضعیت مزایده','نتیجه / وضعیت','مرحله فرایند مزایده','مرحله فرایند',
 'نوع تصمیم','تاریخ تصمیم','شماره نامه','مرجع جلسه','متن تصمیم / مصوبه',
 'نوع انشعاب','نوع محاسبه','مبلغ (ریال)','درصد','توضیح استاندارد',
 'نوع سند','شماره / مرجع سند','شماره / مرجع','وضعیت سند','مسیر / لینک فایل',
 'نوع رویداد','تاریخ رویداد','تاریخ','دسته‌بندی','شرح',
 'نوع ایراد','مقدار فعلی','مقدار استاندارد / مقصد','شرح کنترل','وضعیت',
 'نام ملک مادر','وضعیت تطبیق','نوع مرکز','کاربری اصلی','گروه کاربری','نشانی','ملاحظات',
}
def registry_for(header,filename,sheet):
 if header in MAPPED: key,label,typ,domain,entity=MAPPED[header]; status='MAPPED'
 else:
  digest=hashlib.sha1(f'{sheet}:{header}'.encode()).hexdigest()[:12]; key=f'source.{digest}'; label=header or 'فیلد بدون عنوان (فقط سابقه منبع)'; typ='string'; domain='provenance'; entity='RawCell'; status='MAPPED' if header in OPERATIONAL_HEADINGS else ('REFERENCE_ONLY' if header else 'UNRESOLVED')
 obj,_=CanonicalField.objects.get_or_create(key=key,defaults={'persian_label':label,'aliases':[header],'data_type':typ,'domain':domain,'entity':entity,'meaning':label,'authority':filename,'precedence':5,'searchable':status=='MAPPED','filterable':status=='MAPPED','sortable':status=='MAPPED','default_visible':False,'mapping_status':status})
 if header and header not in obj.aliases:
  obj.aliases=[*obj.aliases,header];obj.save(update_fields=['aliases'])
 return obj,status
def validate_date(raw,source,entity,key,field):
 if raw in (None,''): return ''
 try:return normalize_jalali(raw)
 except (ValueError,TypeError):
  Discrepancy.objects.create(source_file=source,entity_type=entity,entity_key=key,field_key=field,observed_value=text(raw),reason='تاریخ شمسی نامعتبر؛ مقدار خام حفظ شد',severity='HIGH')
  return ''
def rows(ws):
 hr,headers=heading_row(ws)
 if not hr:return [],headers,hr
 out=[]
 for n,row in enumerate(ws.iter_rows(min_row=hr+1,values_only=True),hr+1):
  if any(v not in (None,'') for v in row): out.append((n,dict(zip(headers,row)),row))
 return out,headers,hr
@transaction.atomic
def run(package:Path):
 inventory=inspect_package(package)
 digest=sha(package); batch=ImportBatch.objects.create(source_package=package.name,package_sha256=digest)
 # Authority inputs are immutable evidence. Extraction is confined to runtime data.
 extract=Path(__file__).resolve().parents[1]/'data'/'imports'/digest
 extract.mkdir(parents=True,exist_ok=True)
 with zipfile.ZipFile(package) as z:
  for member in z.infolist():
   try: name=member.filename.encode('cp437').decode('utf-8')
   except (UnicodeEncodeError,UnicodeDecodeError): name=member.filename
   target=extract/Path(name).name
   with z.open(member) as src, target.open('wb') as dst: dst.write(src.read())
 sources={}
 for path in sorted(extract.glob('*.xlsx')):
  sf=SourceFile.objects.create(batch=batch,filename=path.name,sha256=sha(path),byte_size=path.stat().st_size); sources[path.name]=(sf,path)
  wb=load_workbook(path,read_only=True,data_only=True)
  for ws in wb.worksheets:
   data,headers,hr=rows(ws)
   regs=[registry_for(h,path.name,ws.title) for h in headers] if hr else []
   cells=[]
   for n,rawrow in enumerate(ws.iter_rows(values_only=True),1):
    if not any(v not in (None,'') for v in rawrow): continue
    fp=hashlib.sha256(json.dumps([text(v) for v in rawrow],ensure_ascii=False).encode()).hexdigest()
    for col,val in enumerate(rawrow,1):
     if val not in (None,''):
      if hr and n>hr and col<=len(headers): header=headers[col-1]; field,status=regs[col-1]
      else: header=''; field,status=registry_for('',path.name,ws.title)
      cells.append(RawCell(source_file=sf,sheet=ws.title,row_number=n,column_number=col,original_heading=header,raw_value=text(val),value_type=type(val).__name__,row_fingerprint=fp,field=field,mapping_status=status))
   RawCell.objects.bulk_create(cells,batch_size=1000)
 # canonical baseline
 mother_name=next(n for n in sources if n.startswith('املاک_مادر'))
 sf,path=sources[mother_name]; wb=load_workbook(path,read_only=True,data_only=True); data,_,_=rows(wb['املاک مادر'])
 for n,r,_ in data:
  identifier=norm_code(r.get('شناسه ملک مادر')); region=region_obj(r.get('منطقه شهرداری'))
  MotherProperty.objects.create(identifier=identifier,name=text(r.get('نام ملک / مجموعه')),region=region,center_type=text(r.get('نوع مرکز')),primary_usage=text(r.get('کاربری اصلی')),usage_group=text(r.get('گروه کاربری')),area=decimal(r.get('مساحت اعیان (مترمربع)')),address=text(r.get('نشانی')),notes=text(r.get('ملاحظات')),source_row=n)
 status_files=[('فضاهای_تجاری_فعال','ACTIVE'),('فضاهای_از_دور','OUT_OF_CYCLE')]
 for prefix,status in status_files:
  name=next(n for n in sources if n.startswith(prefix)); sf,path=sources[name]; wb=load_workbook(path,read_only=True,data_only=True); data,_,_=rows(wb['فضاها'])
  for n,r,_ in data:
   code=norm_code(r.get('کد فضای تجاری')); reg=region_obj(r.get('منطقه شهرداری')); center=Center.objects.get_or_create(name=text(r.get('نام فضا / مرکز')) or 'نامشخص',region=reg,defaults={'is_special':text(r.get('حوزه سازمانی')) not in ('',text(r.get('منطقه شهرداری')))})[0]
   space=CommercialSpace.objects.create(code=code,name=text(r.get('نام فضا / مرکز')),status=status,region=reg,center=center,organizational_scope=text(r.get('حوزه سازمانی')),asset_type=text(r.get('نوع فضا / دارایی')),area=decimal(r.get('مساحت اعیان (مترمربع)')),address=text(r.get('نشانی فضای تجاری')),physical_details=' | '.join(filter(None,[text(r.get('شرح موقعیت مکانی')),text(r.get('مشخصات فیزیکی و امکانات'))])),current_usage=text(r.get('کاربری اصلی فضای تجاری')),proposed_activity=text(r.get('فعالیت / کاربری پیشنهادی')),activity_group=text(r.get('گروه فعالیت پیشنهادی')),previous_usage=text(r.get('کاربری پیشین')),notes=text(r.get('توضیحات')),source_row=n,source_classification=name)
   CommercialSpaceStatusHistory.objects.create(space=space,new_state=status,effective_date='1405/07/06',reason='طبقه‌بندی اولیه مرجع مصوب ۶ مهر ۱۴۰۵',source_file=sf)
  import_history(wb,sf)
  import_discrepancies(wb,sf)
 # explicit links only
 data,_,_=rows(load_workbook(sources[mother_name][1],read_only=True,data_only=True)['فضاهای مرتبط'])
 for n,r,_ in data:
  code=norm_code(r.get('کد فضای تجاری')); mid=norm_code(r.get('شناسه ملک مادر'))
  try: MotherPropertySpaceLink.objects.create(mother_property=MotherProperty.objects.get(identifier=mid),space=CommercialSpace.objects.get(code=code),evidence=text(r.get('وضعیت تطبیق')) or 'ثبت صریح در شیت فضاهای مرتبط',source_file=sources[mother_name][0],source_sheet='فضاهای مرتبط',source_row=n)
  except (MotherProperty.DoesNotExist,CommercialSpace.DoesNotExist): Discrepancy.objects.create(source_file=sources[mother_name][0],entity_type='MotherPropertySpaceLink',entity_key=f'{mid}:{code}',reason='شناسه یک سوی ارتباط در مرجع canonical یافت نشد',severity='HIGH')
 counts={'mother_properties':MotherProperty.objects.count(),'active_spaces':CommercialSpace.objects.filter(status='ACTIVE').count(),'out_of_cycle_spaces':CommercialSpace.objects.filter(status='OUT_OF_CYCLE').count(),'unique_spaces':CommercialSpace.objects.values('code').distinct().count(),'raw_cells':RawCell.objects.count(),'canonical_fields':CanonicalField.objects.count(),'discrepancies':Discrepancy.objects.count()}
 expected=(inventory.mother_properties,inventory.active_spaces,inventory.out_of_cycle_spaces,inventory.unique_spaces)
 actual=(counts['mother_properties'],counts['active_spaces'],counts['out_of_cycle_spaces'],counts['unique_spaces'])
 if actual!=expected: raise ValueError(f'Canonical import differs from independently inventoried workbooks: actual={actual}, source={expected}')
 batch.status='COMPLETE'; batch.completed_at=timezone.now(); batch.summary=counts; batch.save(); return counts
def region_obj(v):
 val=norm_code(v)
 if not val:return None
 return Region.objects.get_or_create(code=val,defaults={'name':f'منطقه {val}'})[0]
def import_history(wb,sf):
 for sheet,kind in [('بهره‌برداران','beneficiary'),('قراردادها','contract'),('کارشناسی‌ها','appraisal'),('مزایده‌ها','auction'),('مصوبات و دستورات','decision'),('انشعابات و تعهدات','utility'),('اسناد','document'),('سوابق و توضیحات','event')]:
  if sheet not in wb.sheetnames:continue
  data,_,_=rows(wb[sheet])
  for n,r,_ in data:
   code=norm_code(r.get('کد فضای تجاری'))
   try:space=CommercialSpace.objects.get(code=code)
   except CommercialSpace.DoesNotExist:continue
   if kind=='beneficiary':
    name=text(r.get('نام بهره‌بردار / شخص'))
    if name in INVALID_BENEFICIARIES:
     Discrepancy.objects.create(source_file=sf,entity_type='CommercialSpace',entity_key=code,field_key='beneficiary.name',observed_value=name,reason='مقدار وضعیت/فرایند به‌جای هویت بهره‌بردار در منبع ثبت شده است؛ مقدار خام حفظ شد',severity='HIGH')
    elif name:
     b,_=Beneficiary.objects.get_or_create(name=name,defaults={'kind':'UNSPECIFIED'});assignment=BeneficiaryAssignment.objects.create(space=space,beneficiary=b,role=text(r.get('نوع نقش')).replace('Legacy','منبع پیشین'),start_date=validate_date(r.get('تاریخ شروع'),sf,'CommercialSpace',code,'beneficiary.start_date'),end_date=validate_date(r.get('تاریخ پایان'),sf,'CommercialSpace',code,'beneficiary.end_date'),status=text(r.get('وضعیت رابطه')),source_file=sf,source_row=n)
     TimelineEvent.objects.create(space=space,event_type='BENEFICIARY_HISTORY',jalali_date=assignment.start_date,source_entity='BeneficiaryAssignment',source_entity_id=str(assignment.pk),title=f'ثبت سابقه بهره‌بردار: {name}',new_state=assignment.status,provenance=f'{sf.filename} / {sheet} / ردیف {n}')
   elif kind=='contract':
    raw_number=text(r.get('شماره قرارداد'));number='' if raw_number in INVALID_CONTRACT_NUMBERS else raw_number
    signed=normalized_date(r.get('تاریخ انعقاد'));start=normalized_date(r.get('تاریخ شروع'));end=normalized_date(r.get('تاریخ پایان'));amount=decimal(first_value(r,'مبلغ قرارداد با ارزش افزوده (ریال)','مبلغ با ارزش افزوده (ریال)'));status=text(r.get('وضعیت قرارداد'))
    if raw_number and not number:Discrepancy.objects.create(source_file=sf,entity_type='CommercialSpace',entity_key=code,field_key='contract.number',observed_value=raw_number,reason='عبارت وضعیت/فرایند در ستون شماره قرارداد منبع ثبت شده است؛ به‌عنوان شماره قرارداد وارد نشد',severity='HIGH')
    if number or signed or start or end or (amount not in (None,0)) or status:
     contract=Contract.objects.create(space=space,number=number,signed_date=signed,start_date=start,end_date=end,amount_rial=amount,investment_commitment_rial=decimal(r.get('تعهد سرمایه‌گذاری (ریال)')),status=status,signed_state=text(r.get('وضعیت نسخه قرارداد')),source_file=sf,source_row=n)
     TimelineEvent.objects.create(space=space,event_type='CONTRACT_HISTORY',jalali_date=signed or start or end,source_entity='Contract',source_entity_id=str(contract.pk),title=f'سابقه قرارداد {number}' if number else 'سابقه قراردادی بدون شماره معتبر',new_state=status,provenance=f'{sf.filename} / {sheet} / ردیف {n}')
   elif kind=='appraisal' and any(first_value(r,x) not in (None,'') for x in ['سال','مبلغ کارشناسی اجاره ماهانه (ریال)','مبلغ کارشناسی (ریال)']):
    review=text(r.get('وضعیت بررسی'));reported=r.get('تاریخ کارشناسی / گزارش');date=normalized_date(reported) or normalized_date(review);quality='' if normalized_date(review) else review
    if reported not in (None,'') and not normalized_date(reported):Discrepancy.objects.create(source_file=sf,entity_type='CommercialSpace',entity_key=code,field_key='appraisal.appraisal_date',observed_value=text(reported),expected_value=date,reason='مقدار ستون تاریخ کارشناسی نامعتبر است؛ تاریخ معتبر ستون بررسی در صورت وجود استفاده شد',severity='HIGH')
    appraisal=Appraisal.objects.create(space=space,year=text(r.get('سال')),sequence=text(r.get('نوبت')),amount_rial=decimal(first_value(r,'مبلغ کارشناسی اجاره ماهانه (ریال)','مبلغ کارشناسی (ریال)')),appraiser=text(first_value(r,'نام کارشناس ارزیابی','نام کارشناس')),reference=text(first_value(r,'شماره / مرجع نامه کارشناسی','مرجع نامه')),appraisal_date=date,status=quality,source_file=sf,source_row=n)
    TimelineEvent.objects.create(space=space,event_type='APPRAISAL_HISTORY',jalali_date=date,source_entity='Appraisal',source_entity_id=str(appraisal.pk),title='سابقه کارشناسی',description=appraisal.reference,new_state=quality,provenance=f'{sf.filename} / {sheet} / ردیف {n}')
   elif kind=='auction' and any(r.get(x) not in (None,'') for x in ['سال','نتیجه / وضعیت مزایده']):
    item=Auction.objects.create(space=space,year=text(r.get('سال')),sequence=text(r.get('نوبت')),result=text(r.get('نتیجه / وضعیت مزایده')),stage=text(r.get('مرحله فرایند مزایده')),notes=text(r.get('توضیحات')),source_file=sf,source_row=n)
    TimelineEvent.objects.create(space=space,event_type='AUCTION_HISTORY',source_entity='Auction',source_entity_id=str(item.pk),title='سابقه مزایده',description=item.notes,new_state=item.result or item.stage,provenance=f'{sf.filename} / {sheet} / ردیف {n}')
   elif kind=='decision' and text(r.get('متن تصمیم / مصوبه')):
    item=DecisionOrder.objects.create(space=space,decision_type=text(r.get('نوع تصمیم')),decision_date=validate_date(r.get('تاریخ تصمیم'),sf,'CommercialSpace',code,'decision.date'),letter_number=text(r.get('شماره نامه')),session_reference=text(r.get('مرجع جلسه')),text=text(r.get('متن تصمیم / مصوبه')),review_status=text(r.get('وضعیت بررسی')),source_file=sf,source_row=n)
    TimelineEvent.objects.create(space=space,event_type='DECISION_HISTORY',jalali_date=item.decision_date,source_entity='DecisionOrder',source_entity_id=str(item.pk),title=item.decision_type or 'مصوبه یا دستور',description=item.text,new_state=item.review_status,provenance=f'{sf.filename} / {sheet} / ردیف {n}')
   elif kind=='utility' and any(r.get(x) not in (None,'') for x in ['نوع انشعاب','نوع محاسبه','مبلغ (ریال)','درصد','توضیح استاندارد','توضیحات']):
    item=UtilityObligation.objects.create(space=space,utility_type=text(r.get('نوع انشعاب')) or 'نوع در منبع مشخص نشده',calculation_type=text(r.get('نوع محاسبه')),amount_rial=decimal(r.get('مبلغ (ریال)')),percentage=decimal(r.get('درصد')),standard_description=text(r.get('توضیح استاندارد')),notes=text(r.get('توضیحات')),source_file=sf,source_row=n)
    TimelineEvent.objects.create(space=space,event_type='UTILITY_OBLIGATION_HISTORY',source_entity='UtilityObligation',source_entity_id=str(item.pk),title=f'سابقه تعهد انشعاب: {item.utility_type}',description=item.standard_description or item.notes,provenance=f'{sf.filename} / {sheet} / ردیف {n}')
   elif kind=='document' and text(r.get('نوع سند')):
    item=SourceDocumentReference.objects.create(space=space,document_type=text(r.get('نوع سند')),reference=text(r.get('شماره / مرجع سند') or r.get('شماره / مرجع')),status=text(r.get('وضعیت سند')),source_path=text(r.get('مسیر / لینک فایل')),notes=text(r.get('توضیحات')),source_file=sf,source_row=n)
    TimelineEvent.objects.create(space=space,event_type='DOCUMENT_REFERENCE_HISTORY',source_entity='SourceDocumentReference',source_entity_id=str(item.pk),title=f'مرجع سند: {item.document_type}',description=item.reference,new_state=item.status,provenance=f'{sf.filename} / {sheet} / ردیف {n}')
   elif kind=='event':
    date=validate_date(first_value(r,'تاریخ رویداد','تاریخ'),sf,'CommercialSpace',code,'timeline.date')
    TimelineEvent.objects.create(space=space,event_type='SOURCE_HISTORY',jalali_date=date,source_entity='RawSource',source_entity_id=f'{sf.id}:{sheet}:{n}',title=text(r.get('دسته‌بندی')) or 'سابقه ثبت‌شده در منبع',description=text(r.get('شرح')),provenance=f'{sf.filename} / {sheet} / ردیف {n}')

def import_discrepancies(wb,sf):
 sheet='موارد نیازمند بررسی'
 if sheet not in wb.sheetnames:return
 data,_,_=rows(wb[sheet])
 for n,r,_ in data:
  code=norm_code(r.get('کد فضای تجاری'))
  Discrepancy.objects.create(source_file=sf,entity_type='CommercialSpace',entity_key=code,field_key=text(r.get('نوع ایراد')),observed_value=text(r.get('مقدار فعلی')),expected_value=text(r.get('مقدار استاندارد / مقصد')),reason=text(r.get('شرح کنترل')) or 'مورد اعلام‌شده در مرجع مصوب',severity='MEDIUM',status='RESOLVED' if text(r.get('وضعیت')) in ('حل‌شده','تأییدشده') else 'OPEN')
